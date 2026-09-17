"""Read and fill the supported subset of structured DOCX templates."""

from __future__ import annotations

import hashlib
import re
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.table import Table
from docx.text.paragraph import Paragraph

from .models import TemplateField, TemplateSchema, TemplateSection, TemplateTable


PLACEHOLDER = re.compile(r"\{\{\s*([a-zA-Z0-9_\u4e00-\u9fff]+)\s*\}\}|【待填写】")
HEADING = re.compile(r"^Heading ([123])$")


def body_blocks(document):
    for child in document.element.body.iterchildren():
        if child.tag.endswith("}p"):
            yield Paragraph(child, document)
        elif child.tag.endswith("}tbl"):
            yield Table(child, document)


def _replace_paragraph(paragraph: Paragraph, old: str, new: str) -> None:
    if old not in paragraph.text:
        return
    text = paragraph.text.replace(old, new)
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


class TemplateParser:
    def parse(self, path: str | Path) -> TemplateSchema:
        source = Path(path).resolve()
        if source.suffix.lower() != ".docx" or not source.is_file():
            raise ValueError("template must be an existing .docx file")
        document = Document(source)
        sections: list[TemplateSection] = []
        fields: list[TemplateField] = []
        tables: list[TemplateTable] = []
        stack: list[TemplateSection] = []
        for block_index, block in enumerate(body_blocks(document)):
            if isinstance(block, Paragraph):
                match = HEADING.match(block.style.name)
                if match:
                    level = int(match.group(1))
                    while stack and stack[-1].level >= level:
                        stack.pop()
                    section = TemplateSection(
                        section_id=f"s{len(sections)+1:03d}", title=block.text.strip(),
                        level=level, order=len(sections)+1,
                        parent_section_id=stack[-1].section_id if stack else None,
                        heading_location=block_index,
                    )
                    sections.append(section)
                    stack.append(section)
                elif stack:
                    for match in PLACEHOLDER.finditer(block.text):
                        name = match.group(1) or stack[-1].title
                        fields.append(TemplateField(
                            field_id=f"f{len(fields)+1:03d}", field_name=name,
                            field_type="paragraph", required=True,
                            placeholder=match.group(0),
                            location={"block": block_index}, section_id=stack[-1].section_id,
                        ))
            elif isinstance(block, Table):
                table = TemplateTable(table_id=f"t{len(tables)+1:03d}", rows=len(block.rows))
                if stack:
                    for row_index, row in enumerate(block.rows):
                        for cell_index, cell in enumerate(row.cells):
                            for match in PLACEHOLDER.finditer(cell.text):
                                label = row.cells[cell_index-1].text.strip() if cell_index else ""
                                name = match.group(1) or label or stack[-1].title
                                location = {"block": block_index, "row": row_index, "cell": cell_index}
                                table.fillable_cells.append(location)
                                fields.append(TemplateField(
                                    field_id=f"f{len(fields)+1:03d}", field_name=name,
                                    field_type="table", required=True,
                                    placeholder=match.group(0), location=location,
                                    section_id=stack[-1].section_id,
                                ))
                tables.append(table)
        if not sections:
            raise ValueError("template has no Heading 1/2/3 sections")
        explicit = {field.section_id for field in fields}
        for section in sections:
            if section.section_id not in explicit and not any(
                child.parent_section_id == section.section_id for child in sections
            ):
                fields.append(TemplateField(
                    field_id=f"f{len(fields)+1:03d}", field_name=section.title,
                    field_type="implicit", required=True, placeholder=None,
                    location={"block": section.heading_location}, section_id=section.section_id,
                ))
        digest = hashlib.sha256(source.read_bytes()).hexdigest()[:16]
        return TemplateSchema(f"tpl_{digest}", source.name, sections, fields, tables)


class TemplateRenderer:
    def render(self, template: str | Path, output: str | Path, schema: TemplateSchema, drafts) -> Path:
        source, target = Path(template).resolve(), Path(output).resolve()
        if source == target:
            raise ValueError("draft output must differ from original template")
        document = Document(source)
        blocks = list(body_blocks(document))
        values = {field_id: value for draft in drafts for field_id, value in draft.field_values.items()}
        for field in schema.fields:
            value = values.get(field.field_id, "[MISSING: 当前知识库中未发现明确资料]")
            block = blocks[field.location["block"]]
            if field.field_type == "table":
                cell = block.cell(field.location["row"], field.location["cell"])
                for paragraph in cell.paragraphs:
                    _replace_paragraph(paragraph, field.placeholder or "", value)
            elif field.field_type == "paragraph":
                _replace_paragraph(block, field.placeholder or "", value)
            else:
                element = OxmlElement("w:p")
                block._p.addnext(element)
                Paragraph(element, block._parent).add_run(value)
        target.parent.mkdir(parents=True, exist_ok=True)
        document.save(target)
        return target
