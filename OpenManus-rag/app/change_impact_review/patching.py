"""Fail-closed paragraph patching with one deterministic Quality Gate result."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document

from app.document_workflow.evidence_models import Evidence, sha256_text
from app.document_workflow.evidence_selection import evidence_in_scope
from app.document_workflow.scope import WorkflowScope
from app.document_workflow.freshness import EvidenceFreshness, EvidenceFreshnessValidator

from .models import (
    PatchApplyResult,
    PatchApplyStatus,
    PatchCandidate,
    PatchReviewStatus,
    QualityGateReason,
    QualityGateResult,
    QualityGateStatus,
)


class ParagraphPatchApplier:
    def evaluate(
        self,
        patch: PatchCandidate,
        *,
        source: str | Path,
        output: str | Path,
        current_version_id: str,
        applied_keys: set[str],
    ) -> QualityGateResult:
        gate, _, _ = self._inspect(
            patch,
            source=source,
            output=output,
            current_version_id=current_version_id,
            applied_keys=applied_keys,
        )
        return gate

    def apply(
        self,
        patch: PatchCandidate,
        *,
        source: str | Path,
        output: str | Path,
        current_version_id: str,
        applied_keys: set[str],
    ) -> PatchApplyResult:
        gate, document, index = self._inspect(
            patch,
            source=source,
            output=output,
            current_version_id=current_version_id,
            applied_keys=applied_keys,
        )
        if gate.status is QualityGateStatus.BLOCKED:
            patch.apply_status = gate.patch_apply_status
            output_path = (
                str(Path(output).resolve())
                if gate.reasons == [QualityGateReason.ALREADY_APPLIED] and Path(output).exists()
                else None
            )
            return PatchApplyResult(
                patch_id=patch.patch_id,
                status=patch.apply_status,
                output_path=output_path,
                failure_reason=gate.failure_reason,
                idempotency_key=patch.idempotency_key,
                quality_gate=gate,
            )

        assert document is not None and index is not None
        paragraph = document.paragraphs[index]
        if paragraph.runs:
            paragraph.runs[0].text = patch.proposed_content
            for run in paragraph.runs[1:]:
                run.text = ""
        else:
            paragraph.add_run(patch.proposed_content)
        output_path = Path(output).resolve()
        output_path.parent.mkdir(parents=True, exist_ok=True)
        document.save(output_path)
        applied_keys.add(patch.idempotency_key)
        patch.apply_status = PatchApplyStatus.APPLIED
        return PatchApplyResult(
            patch_id=patch.patch_id,
            status=patch.apply_status,
            output_path=str(output_path),
            idempotency_key=patch.idempotency_key,
            quality_gate=QualityGateResult.passed(PatchApplyStatus.APPLIED),
        )

    def _inspect(
        self,
        patch: PatchCandidate,
        *,
        source: str | Path,
        output: str | Path,
        current_version_id: str,
        applied_keys: set[str],
    ) -> tuple[QualityGateResult, Document | None, int | None]:
        key = patch.idempotency_key
        if key in applied_keys:
            return (
                QualityGateResult.blocked(
                    QualityGateReason.ALREADY_APPLIED,
                    PatchApplyStatus.SKIPPED_ALREADY_APPLIED,
                ),
                None,
                None,
            )
        if patch.review_status is not PatchReviewStatus.APPROVED:
            return (
                QualityGateResult.blocked(
                    QualityGateReason.REVIEW_REQUIRED,
                    PatchApplyStatus.BLOCKED_REVIEW,
                    failure_reason="PATCH_NOT_APPROVED",
                ),
                None,
                None,
            )
        if current_version_id != patch.base_version_id:
            return (
                QualityGateResult.blocked(
                    QualityGateReason.BASE_VERSION_MISMATCH,
                    PatchApplyStatus.CONFLICT,
                    failure_reason="BASE_VERSION_CHANGED",
                ),
                None,
                None,
            )
        match = re.fullmatch(r"paragraph:(\d+)", patch.target_anchor)
        if match is None:
            return (
                QualityGateResult.blocked(
                    QualityGateReason.TARGET_CHANGED,
                    PatchApplyStatus.CONFLICT,
                    failure_reason="TARGET_ANCHOR_INVALID",
                ),
                None,
                None,
            )
        source_path = Path(source).resolve()
        output_path = Path(output).resolve()
        if source_path == output_path:
            raise ValueError("candidate output must not overwrite its source")
        if not source_path.is_file() or source_path.suffix.lower() != ".docx":
            raise ValueError("patch source must be an existing DOCX")
        document = Document(source_path)
        index = int(match.group(1))
        if index >= len(document.paragraphs):
            return (
                QualityGateResult.blocked(
                    QualityGateReason.TARGET_CHANGED,
                    PatchApplyStatus.CONFLICT,
                    failure_reason="TARGET_ANCHOR_MISSING",
                ),
                None,
                None,
            )
        paragraph = document.paragraphs[index]
        if (
            paragraph.text != patch.original_content
            or sha256_text(paragraph.text) != patch.original_content_hash
        ):
            return (
                QualityGateResult.blocked(
                    QualityGateReason.TARGET_CHANGED,
                    PatchApplyStatus.CONFLICT,
                    failure_reason="ORIGINAL_CONTENT_CHANGED",
                ),
                None,
                None,
            )
        if output_path.exists():
            return (
                QualityGateResult.blocked(
                    QualityGateReason.TARGET_CHANGED,
                    PatchApplyStatus.CONFLICT,
                    failure_reason="CANDIDATE_OUTPUT_EXISTS",
                ),
                None,
                None,
            )
        return QualityGateResult.passed(), document, index


class PatchExecutionService:
    def __init__(self, active_version_for_document):
        self.freshness = EvidenceFreshnessValidator(active_version_for_document)
        self.applier = ParagraphPatchApplier()

    def evaluate(
        self,
        patch: PatchCandidate,
        *,
        evidence_by_id: dict[str, Evidence],
        expected_scope_fingerprint: str,
        source: str | Path,
        output: str | Path,
        current_version_id: str,
        applied_keys: set[str],
        expected_scope: dict | WorkflowScope | None = None,
    ) -> QualityGateResult:
        evidence_gate = self._evaluate_evidence(
            patch,
            evidence_by_id=evidence_by_id,
            expected_scope_fingerprint=expected_scope_fingerprint,
            expected_scope=expected_scope,
        )
        if evidence_gate.status is QualityGateStatus.BLOCKED:
            return evidence_gate
        return self.applier.evaluate(
            patch,
            source=source,
            output=output,
            current_version_id=current_version_id,
            applied_keys=applied_keys,
        )

    def apply(
        self,
        patch: PatchCandidate,
        *,
        evidence_by_id: dict[str, Evidence],
        expected_scope_fingerprint: str,
        source: str | Path,
        output: str | Path,
        current_version_id: str,
        applied_keys: set[str],
        expected_scope: dict | WorkflowScope | None = None,
    ) -> PatchApplyResult:
        evidence_gate = self._evaluate_evidence(
            patch,
            evidence_by_id=evidence_by_id,
            expected_scope_fingerprint=expected_scope_fingerprint,
            expected_scope=expected_scope,
        )
        if evidence_gate.status is QualityGateStatus.BLOCKED:
            patch.apply_status = evidence_gate.patch_apply_status
            return PatchApplyResult(
                patch_id=patch.patch_id,
                status=patch.apply_status,
                failure_reason=evidence_gate.failure_reason,
                idempotency_key=patch.idempotency_key,
                quality_gate=evidence_gate,
            )
        return self.applier.apply(
            patch,
            source=source,
            output=output,
            current_version_id=current_version_id,
            applied_keys=applied_keys,
        )

    def _evaluate_evidence(
        self,
        patch: PatchCandidate,
        *,
        evidence_by_id: dict[str, Evidence],
        expected_scope_fingerprint: str,
        expected_scope: dict | WorkflowScope | None = None,
    ) -> QualityGateResult:
        if patch.scope_fingerprint != expected_scope_fingerprint:
            return QualityGateResult.blocked(
                QualityGateReason.OUT_OF_SCOPE_EVIDENCE,
                PatchApplyStatus.BLOCKED_SCOPE,
                failure_reason="SCOPE_FINGERPRINT_CHANGED",
            )
        for evidence_id in patch.evidence_ids:
            item = evidence_by_id.get(evidence_id)
            if item is None:
                return QualityGateResult.blocked(
                    QualityGateReason.NO_EVIDENCE,
                    PatchApplyStatus.BLOCKED_STALE_EVIDENCE,
                    failure_reason="EVIDENCE_MISSING",
                )
            scope = (
                expected_scope
                if isinstance(expected_scope, WorkflowScope)
                else WorkflowScope.from_dict(expected_scope)
                if expected_scope is not None
                else None
            )
            if scope is not None and not evidence_in_scope(item, scope):
                return QualityGateResult.blocked(
                    QualityGateReason.OUT_OF_SCOPE_EVIDENCE,
                    PatchApplyStatus.BLOCKED_SCOPE,
                    failure_reason="EVIDENCE_OUT_OF_SCOPE",
                )
            if item.content_hash != patch.evidence_content_hashes.get(evidence_id):
                return QualityGateResult.blocked(
                    QualityGateReason.STALE_EVIDENCE,
                    PatchApplyStatus.BLOCKED_STALE_EVIDENCE,
                    failure_reason="EVIDENCE_CONTENT_CHANGED",
                )
            decision = self.freshness.evaluate(item)
            if decision.freshness is not EvidenceFreshness.FRESH:
                return QualityGateResult.blocked(
                    QualityGateReason.STALE_EVIDENCE,
                    PatchApplyStatus.BLOCKED_STALE_EVIDENCE,
                    failure_reason="EVIDENCE_STALE",
                )
        return QualityGateResult.passed()
