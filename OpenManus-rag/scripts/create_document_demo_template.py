"""Create a small structured, simulated project-plan DOCX template."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml.ns import qn


def create(path: Path) -> None:
    document = Document()
    document.sections[0].page_width = Inches(8.5)
    document.sections[0].page_height = Inches(11)
    styles = document.styles
    styles["Normal"].font.size = Pt(11)
    for name in ("Title", "Heading 1", "Heading 2", "Heading 3"):
        styles[name].font.color.rgb = RGBColor(0, 0, 0)
    title_props = styles["Title"]._element.pPr
    if title_props is not None:
        border = title_props.find(qn("w:pBdr"))
        if border is not None:
            title_props.remove(border)
    document.add_paragraph("研发项目实施计划", style="Title")
    document.add_paragraph("本文件用于整理模拟项目资料。生成内容均需人工审核。")
    document.add_heading("1 项目概述", level=1)
    document.add_heading("1.1 项目背景", level=2)
    document.add_paragraph("{{项目背景}}")
    document.add_heading("1.2 项目目标", level=2)
    document.add_paragraph("{{项目目标}}")
    document.add_heading("2 实施安排", level=1)
    document.add_heading("2.1 阶段划分", level=2)
    document.add_paragraph("{{阶段划分}}")
    document.add_heading("2.2 性能要求", level=2)
    table = document.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.cell(0, 0).text = "吞吐能力"
    table.cell(0, 1).text = "【待填写】"
    document.add_heading("3 验收", level=1)
    document.add_heading("3.1 验收依据", level=2)
    document.add_paragraph("{{验收依据}}")
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(path)


if __name__ == "__main__":
    create(Path("project_delivery/document_workflow_v2/demo/template.docx"))
