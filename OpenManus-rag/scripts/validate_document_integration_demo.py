"""Read-only checks of the synthetic HTTP document-workflow demo outputs."""

from __future__ import annotations

import argparse
import hashlib
import json
import re
from pathlib import Path

from docx import Document


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("demo_dir", type=Path)
    output = parser.parse_args().demo_dir
    template_path = output / "template.docx"
    draft_path = output / "draft.docx"
    evidence = json.loads((output / "evidence.json").read_text(encoding="utf-8"))
    trace = json.loads((output / "execution_trace.json").read_text(encoding="utf-8"))
    assert trace["rag_adapter"] == "HTTPRetrieveClient"
    assert trace["retrieval_endpoint"] == "/retrieve"
    assert trace["total_rag_calls"] >= 5
    assert len(trace["sections"]) == 5
    assert sum(section["rag_calls"] for section in trace["sections"]) == trace["total_rag_calls"]
    assert sum(section["reused_evidence_count"] for section in trace["sections"]) > 0
    assert sum(section["new_evidence_count"] for section in trace["sections"]) == trace["total_unique_evidence"]
    assert any("吞吐能力" in section["missing_fields"] for section in trace["sections"])
    assert trace["requires_human_review"] is evidence["requires_human_review"] is True
    assert hashlib.sha256(template_path.read_bytes()).hexdigest() == trace["original_template_hash"]
    assert draft_path.is_file() and draft_path != template_path
    original = Document(template_path)
    draft = Document(draft_path)
    headings = lambda doc: [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
    assert headings(original) == headings(draft)
    assert len(original.tables) == len(draft.tables) == 1
    assert "[MISSING:" in draft.tables[0].cell(0, 1).text
    allowed = {item["evidence_id"] for item in evidence["evidence"]}
    assert len(allowed) == trace["total_unique_evidence"]
    text = "\n".join(p.text for p in draft.paragraphs)
    text += "\n" + "\n".join(cell.text for table in draft.tables for row in table.rows for cell in row.cells)
    cited = set(re.findall(r"\[Evidence: (ev_[0-9a-f]{20})\]", text))
    assert len(cited) >= 4 and cited.issubset(allowed)
    print("SAFE_END_TO_END_DEMO=PASS MULTI_SECTION_REAL_RAG_CALLS=PASS")
    print("EVIDENCE_CACHE=PASS EVIDENCE_DEDUP=PASS EVIDENCE_MEMBERSHIP=PASS")
    print("MISSING_FIELD=PASS DOCX_STRUCTURE=PASS ORIGINAL_TEMPLATE_UNCHANGED=PASS")
    print(f"rag_calls={trace['total_rag_calls']} unique_evidence={len(allowed)} citations={len(cited)}")


if __name__ == "__main__":
    main()
