from __future__ import annotations

import hashlib
from datetime import datetime, timezone

import httpx
import pytest
from docx import Document

from app.document_workflow.evidence import EvidenceCache
from app.document_workflow.freshness import EvidenceFreshness, EvidenceFreshnessValidator
from app.document_workflow.rag import HTTPRetrieveClient, RAGTool
from app.document_workflow.state import CheckpointStore
from app.document_workflow.workflow import DocumentWorkflow
from app.research.models import Evidence, SourceLevel, normalize_text


def evidence(*, document="REQ", version="REQ@1", chunk="REQ@1:c1", freshness="FRESH"):
    return Evidence(
        title=document,
        content="最大并发为 500。",
        organization="RAG",
        source_type="RAG",
        document_number=document,
        document_id=document,
        version_id=version,
        version_status="ACTIVE" if freshness == "FRESH" else "SUPERSEDED",
        freshness=freshness,
        chunk_id=chunk,
        section="容量",
        page_number=1,
        source_level=SourceLevel.TIER3,
        retrieved_at=datetime.now(timezone.utc),
    )


def test_fresh_stale_unknown_and_explicit_historical_decisions():
    active = {"REQ": "REQ@2"}
    validator = EvidenceFreshnessValidator(active.get)
    fresh = evidence(version="REQ@2", chunk="REQ@2:c1")
    stale = evidence(version="REQ@1", chunk="REQ@1:c1")
    unknown = evidence(document="UNKNOWN", version="UNKNOWN@1", chunk="UNKNOWN@1:c1")
    assert validator.evaluate(fresh).freshness == EvidenceFreshness.FRESH
    assert validator.evaluate(stale).freshness == EvidenceFreshness.STALE
    assert validator.evaluate(unknown).freshness == EvidenceFreshness.UNKNOWN
    assert validator.evaluate(
        stale, historical_version_ids={"REQ@1"}
    ).freshness == EvidenceFreshness.FRESH


def test_version_and_chunk_participate_in_stable_evidence_identity():
    first = evidence(version="REQ@1", chunk="REQ@1:c1")
    same = evidence(version="REQ@1", chunk="REQ@1:c1")
    newer = evidence(version="REQ@2", chunk="REQ@2:c1")
    assert first.evidence_id == same.evidence_id
    assert first.evidence_id != newer.evidence_id


def test_cache_marks_only_tasks_using_stale_evidence_for_refresh(tmp_path):
    cache = EvidenceCache(tmp_path, tmp_path / "evidence.json")
    stale = evidence(version="REQ@1", chunk="REQ@1:c1")
    fresh = evidence(document="DESIGN", version="DESIGN@1", chunk="DESIGN@1:c1")
    cache.add("task-stale", "并发", stale)
    cache.add("task-fresh", "设计", fresh)
    decisions, affected = cache.apply_freshness(
        EvidenceFreshnessValidator({"REQ": "REQ@2", "DESIGN": "DESIGN@1"}.get)
    )
    assert affected == {"task-stale"}
    assert {item.freshness for item in decisions} == {
        EvidenceFreshness.STALE, EvidenceFreshness.FRESH,
    }
    assert cache.for_task("task-stale") == []
    assert not cache.membership("task-stale", [stale.evidence_id])
    assert cache.membership("task-fresh", [fresh.evidence_id])


def _three_section_template(path):
    document = Document()
    for title in ("项目背景", "项目目标", "阶段划分"):
        document.add_heading(title, level=1)
        document.add_paragraph("{{" + title + "}}")
    document.save(path)


def test_resume_detects_stale_and_reretrieves_only_affected_completed_task(monkeypatch, tmp_path):
    source = tmp_path / "template.docx"
    output = tmp_path / "output"
    _three_section_template(source)
    state = {"interrupt": True, "req_version": "REQ@1", "post_queries": []}

    def fake_post(url, *, json, timeout):
        query = json["query"]
        state["post_queries"].append(query)
        if "项目背景" in query:
            document_id, version_id, content = "REQ", state["req_version"], f"项目背景：{state['req_version']}。"
        elif "项目目标" in query:
            document_id, version_id, content = "DESIGN", "DESIGN@1", "项目目标：完成设计。"
        else:
            document_id, version_id, content = "PLAN", "PLAN@1", "阶段划分：开发与验收。"
        raw = normalize_text(content).encode("utf-8")
        hit = {
            "chunk_id": f"{version_id}:c1",
            "document_id": document_id,
            "version_id": version_id,
            "version_status": "ACTIVE",
            "section_id": f"{document_id}:sec:1",
            "section_path": [query],
            "page_number": 1,
            "content": content,
            "content_hash": hashlib.sha256(raw).hexdigest(),
            "similarity": 1.0,
            "rank": 1,
        }
        return httpx.Response(200, json={"query": query, "results": [hit]},
                              request=httpx.Request("POST", url))

    def fake_get(url, *, timeout):
        payload = {"documents": [
            {"document_id": "REQ", "active_version": {"version_id": "REQ@2"}},
            {"document_id": "DESIGN", "active_version": {"version_id": "DESIGN@1"}},
            {"document_id": "PLAN", "active_version": {"version_id": "PLAN@1"}},
        ]}
        return httpx.Response(200, json=payload, request=httpx.Request("GET", url))

    original = RAGTool.retrieve_knowledge

    def interrupt_third(self, query, top_k=None):
        if state["interrupt"] and "阶段划分" in query:
            raise KeyboardInterrupt("checkpoint after two completed tasks")
        return original(self, query, top_k)

    monkeypatch.setattr(httpx, "post", fake_post)
    monkeypatch.setattr(httpx, "get", fake_get)
    monkeypatch.setattr(RAGTool, "retrieve_knowledge", interrupt_third)
    with pytest.raises(KeyboardInterrupt):
        DocumentWorkflow(HTTPRetrieveClient("http://localhost:8000")).run(source, output)
    checkpoint_file = next((output / "checkpoints").glob("dw_*.json"))
    workflow_id = checkpoint_file.stem
    checkpoint = CheckpointStore(output / "checkpoints").load(workflow_id)
    assert checkpoint.section_states["s001"]["status"] == "COMPLETE"
    assert checkpoint.section_states["s002"]["status"] == "COMPLETE"

    state.update(interrupt=False, req_version="REQ@2", post_queries=[])
    result = DocumentWorkflow(HTTPRetrieveClient("http://localhost:8000")).resume(
        workflow_id, output=output
    )
    assert any("项目背景" in query for query in state["post_queries"])
    assert any("阶段划分" in query for query in state["post_queries"])
    assert not any("项目目标" in query for query in state["post_queries"])
    assert result["state"].section_states["s001"]["attempt_count"] == 2
    assert result["state"].section_states["s002"]["attempt_count"] == 1
    assert any(item["freshness"] == "STALE" for item in result["trace"]["evidence_freshness"])
    snapshots = result["state"].evidence_snapshot["evidence"]
    assert any(item.get("version_id") == "REQ@1" and item.get("freshness") == "STALE" for item in snapshots)
    assert any(item.get("version_id") == "REQ@2" and item.get("freshness") == "FRESH" for item in snapshots)


