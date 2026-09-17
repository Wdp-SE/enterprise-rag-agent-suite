from __future__ import annotations

import hashlib
import json

import pytest
import httpx
from docx import Document

from app.document_workflow.evidence import EvidenceCache
from app.document_workflow.planning import SectionTaskPlanner
from app.document_workflow.policy import DOCUMENT_WORKFLOW_PROFILE
from app.document_workflow.rag import DemoRAGClient, HTTPRAGClient, HTTPRetrieveClient, RAGTool
from app.document_workflow.template import TemplateParser, TemplateRenderer
from app.document_workflow.workflow import DocumentWorkflow
from app.reliability.policy import CapabilityId, PolicyContext, PolicyEvaluationPoint, TaskPolicy
from app.reliability.retry import SideEffectLevel
from scripts.create_document_demo_template import create


@pytest.fixture
def template(tmp_path):
    path = tmp_path / "template.docx"
    create(path)
    return path


def test_heading_table_schema_and_planning(template):
    schema = TemplateParser().parse(template)
    assert [section.level for section in schema.sections] == [1, 2, 2, 1, 2, 2, 1, 2]
    assert schema.sections[1].parent_section_id == schema.sections[0].section_id
    assert len(schema.tables) == 1
    assert len(schema.tables[0].fillable_cells) == 1
    assert schema.tables[0].fillable_cells[0]["row"] == 0
    assert schema.tables[0].fillable_cells[0]["cell"] == 1
    assert {field.field_name for field in schema.fields} == {
        "项目背景", "项目目标", "阶段划分", "吞吐能力", "验收依据",
    }
    assert len(SectionTaskPlanner().plan(schema)) == 5


def test_rag_adapter_dedup_and_membership(tmp_path):
    tool = RAGTool(DemoRAGClient())
    first = tool.retrieve_knowledge("项目背景")
    assert len(first) == 1
    assert first[0].source_type == "RAG" and first[0].page_number == 1
    cache = EvidenceCache(tmp_path, tmp_path / "store.json")
    assert cache.add("task-1", "项目背景", first[0])
    assert not cache.add("task-1", "项目计划 项目背景", first[0])
    assert len(cache.store) == 1
    assert cache.membership("task-1", [first[0].evidence_id])
    assert not cache.membership("task-1", ["ev_nonexistent"])
    assert not cache.membership("other-task", [first[0].evidence_id])


def test_http_rag_v2_query_contract(monkeypatch):
    def fake_post(url, *, json, timeout):
        assert url == "http://localhost:8000/query"
        assert json == {"question": "项目目标"}
        return httpx.Response(200, json={
            "answer": "项目目标：形成可审核草稿。",
            "sources": [{"document_id": "demo-a", "page_number": 2}],
            "status": "OK", "trace": [],
        }, request=httpx.Request("POST", url))

    monkeypatch.setattr(httpx, "post", fake_post)
    hits = RAGTool(HTTPRAGClient("http://localhost:8000")).retrieve_knowledge("项目目标")
    assert len(hits) == 1
    assert hits[0].document_number == "demo-a"
    assert hits[0].page_number == 2
    assert "形成可审核草稿" in hits[0].content


def test_http_retrieve_metadata_stable_identity_and_dedup(monkeypatch, tmp_path):
    content = "项目背景: 合成研发项目。"
    raw_hit = {
        "chunk_id": "safe-doc:7", "document_id": "safe-doc",
        "section_id": "safe-doc:sec:1", "section_path": ["概述", "项目背景"],
        "page_number": 7, "content": content,
        "content_hash": hashlib.sha256(content.encode()).hexdigest(),
        "similarity": 0.82, "rank": 1,
    }

    def fake_post(url, *, json, timeout):
        assert url == "http://localhost:8000/retrieve"
        assert json["top_k"] == 3
        return httpx.Response(200, json={"query": json["query"], "results": [raw_hit]},
                              request=httpx.Request("POST", url))

    monkeypatch.setattr(httpx, "post", fake_post)
    tool = RAGTool(HTTPRetrieveClient("http://localhost:8000"))
    first = tool.retrieve_knowledge("项目背景", top_k=3)[0]
    second = tool.retrieve_knowledge("项目计划 项目背景", top_k=3)[0]
    assert first.evidence_id == second.evidence_id
    assert first.chunk_id == raw_hit["chunk_id"]
    assert first.document_number == raw_hit["document_id"]
    assert first.section == raw_hit["section_id"]
    assert first.section_path == raw_hit["section_path"]
    assert first.page_number == raw_hit["page_number"]
    assert first.content_hash == raw_hit["content_hash"]
    cache = EvidenceCache(tmp_path, tmp_path / "store.json")
    assert cache.add("task-1", "项目背景", first)
    assert not cache.add("task-2", "项目计划 项目背景", second)
    assert len(cache.store) == 1
    assert cache.membership("task-2", [first.evidence_id])
    assert not cache.membership("task-2", ["ev_fabricated"])


@pytest.mark.parametrize("failure", ["connection", "timeout", "server", "not_ready", "lifecycle_invalid", "malformed"])
def test_http_retrieve_fails_boundedly(monkeypatch, template, tmp_path, failure):
    def fake_post(url, *, json, timeout):
        if failure == "connection":
            raise httpx.ConnectError("connection refused")
        if failure == "timeout":
            raise httpx.TimeoutException("timed out")
        if failure in {"server", "not_ready", "lifecycle_invalid"}:
            code = 500 if failure == "server" else 503
            return httpx.Response(code, request=httpx.Request("POST", url))
        return httpx.Response(200, json={"query": json["query"], "results": [{}]},
                              request=httpx.Request("POST", url))

    monkeypatch.setattr(httpx, "post", fake_post)
    result = DocumentWorkflow(HTTPRetrieveClient("http://localhost:8000")).run(
        template, tmp_path / "failure"
    )
    assert result["trace"]["errors"]
    assert result["trace"]["total_rag_calls"] <= 5
    assert all(section["status"] == "PARTIAL" for section in result["trace"]["sections"])
    assert result["trace"]["requires_human_review"] is True


def test_heading_3_without_placeholder_creates_task(tmp_path):
    path = tmp_path / "implicit.docx"
    document = Document()
    document.add_heading("项目安排", level=1)
    document.add_heading("实施", level=2)
    document.add_heading("阶段责任", level=3)
    document.save(path)
    schema = TemplateParser().parse(path)
    tasks = SectionTaskPlanner().plan(schema)
    assert [section.level for section in schema.sections] == [1, 2, 3]
    assert len(tasks) == 1
    assert tasks[0].required_fields[0].field_type == "implicit"


def test_implicit_heading_is_rendered_without_changing_source(tmp_path):
    source = tmp_path / "implicit.docx"
    document = Document()
    document.add_heading("项目背景", level=1)
    document.save(source)
    before = hashlib.sha256(source.read_bytes()).hexdigest()
    output = tmp_path / "result"
    DocumentWorkflow(DemoRAGClient()).run(source, output)
    paragraphs = [p.text for p in Document(output / "draft.docx").paragraphs]
    assert paragraphs[0] == "项目背景"
    assert "[Evidence: ev_" in paragraphs[1]
    assert hashlib.sha256(source.read_bytes()).hexdigest() == before


def test_workflow_missing_no_progress_and_docx(template, tmp_path):
    before = hashlib.sha256(template.read_bytes()).hexdigest()
    client = DemoRAGClient()
    output = tmp_path / "output"
    result = DocumentWorkflow(client).run(template, output)
    assert result["trace"]["total_rag_calls"] >= 7
    assert client.calls == result["trace"]["total_rag_calls"]
    assert any(section["status"] == "NO_PROGRESS" for section in result["trace"]["sections"])
    assert any("吞吐能力" in section["missing_fields"] for section in result["trace"]["sections"])
    assert result["trace"]["requires_human_review"] is True
    assert hashlib.sha256(template.read_bytes()).hexdigest() == before
    assert (output / "draft.docx").is_file()
    assert (output / "evidence.json").is_file()
    assert (output / "execution_trace.json").is_file()
    text = "\n".join(p.text for p in Document(output / "draft.docx").paragraphs)
    assert "[Evidence: ev_" in text
    assert "【待填写】" not in Document(output / "draft.docx").tables[0].cell(0, 1).text
    assert "[MISSING:" in Document(output / "draft.docx").tables[0].cell(0, 1).text
    evidence = json.loads((output / "evidence.json").read_text(encoding="utf-8"))
    assert len(evidence["evidence"]) == 4
    allowed = {item["evidence_id"] for item in evidence["evidence"]}
    assert all(set(draft.evidence_ids).issubset(allowed) for draft in result["drafts"])


def test_original_cannot_be_renderer_target(template):
    schema = TemplateParser().parse(template)
    with pytest.raises(ValueError, match="must differ"):
        TemplateRenderer().render(template, template, schema, [])


@pytest.mark.parametrize("capability", [CapabilityId.WEB_SEARCH, CapabilityId.BROWSER_ACQUIRE])
def test_document_policy_denies_external_fact_tools(capability):
    decision = TaskPolicy(DOCUMENT_WORKFLOW_PROFILE).evaluate(PolicyContext(
        workflow="document_workflow", evaluation_point=PolicyEvaluationPoint.PRE_OPERATION,
        operation=capability.value, tool_name=capability.value,
        side_effect_level=SideEffectLevel.READ_ONLY,
        network_required=True, external_write=False,
    ))
    assert not decision.allowed
