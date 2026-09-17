from __future__ import annotations

import hashlib
import json
from dataclasses import replace

import httpx
import pytest
from docx import Document

from app.document_workflow.configuration import DocumentWorkflowConfig
from app.document_workflow.integrity import OutputIntegrityError, OutputIntegrityValidator
from app.document_workflow.rag import DemoRAGClient, HTTPRetrieveClient, RAGTool
from app.document_workflow.state import CheckpointInvalid, CheckpointStore
from app.document_workflow.workflow import DocumentWorkflow


def two_section_template(path):
    document = Document()
    document.add_heading("项目背景", level=1)
    document.add_paragraph("{{项目背景}}")
    document.add_heading("项目目标", level=1)
    document.add_paragraph("{{项目目标}}")
    document.save(path)


@pytest.mark.parametrize("change", [
    {"rag_base_url": "ftp://wrong"}, {"rag_timeout_seconds": 0},
    {"rag_retry_limit": -1}, {"rag_top_k": 0},
    {"no_progress_threshold": 0}, {"execution_budget": 0},
    {"requires_human_review": False},
])
def test_config_fails_fast(change):
    with pytest.raises(ValueError):
        replace(DocumentWorkflowConfig(), **change).validate()


def test_checkpoint_resume_preserves_evidence_and_skips_complete(monkeypatch, tmp_path):
    source = tmp_path / "template.docx"
    output = tmp_path / "output"
    two_section_template(source)
    before = hashlib.sha256(source.read_bytes()).hexdigest()
    original = RAGTool.retrieve_knowledge

    def interrupted(self, query, top_k=None):
        if "项目目标" in query:
            raise KeyboardInterrupt("simulated process stop")
        return original(self, query, top_k)

    with monkeypatch.context() as scoped:
        scoped.setattr(RAGTool, "retrieve_knowledge", interrupted)
        with pytest.raises(KeyboardInterrupt):
            DocumentWorkflow(DemoRAGClient()).run(source, output)

    files = list((output / "checkpoints").glob("dw_*.json"))
    assert len(files) == 1
    workflow_id = files[0].stem
    checkpoint = CheckpointStore(output / "checkpoints")
    state = checkpoint.load(workflow_id)
    assert state.section_states["s001"]["status"] == "COMPLETE"
    first_id = state.section_drafts["s001"]["evidence_ids"][0]
    assert any(record["evidence_id"] == first_id for record in state.evidence_snapshot["evidence"])
    client = DemoRAGClient()
    result = DocumentWorkflow(client).resume(workflow_id, output=output)
    assert client.calls == 1
    assert result["state"].workflow_status == "COMPLETE"
    assert result["drafts"][0].evidence_ids == [first_id]
    assert result["draft_path"].is_file()
    assert hashlib.sha256(source.read_bytes()).hexdigest() == before
    assert "示例研发项目" not in (output / "execution_trace.json").read_text(encoding="utf-8")
    assert "示例研发项目" in files[0].read_text(encoding="utf-8")
    with pytest.raises(CheckpointInvalid, match="completed"):
        DocumentWorkflow(DemoRAGClient()).resume(workflow_id, output=output)


def test_checkpoint_corruption_version_and_template_hash_fail_closed(tmp_path):
    source = tmp_path / "template.docx"
    output = tmp_path / "output"
    two_section_template(source)
    result = DocumentWorkflow(DemoRAGClient()).run(source, output)
    store = CheckpointStore(output / "checkpoints")
    target = store.path(result["state"].workflow_id)
    raw = target.read_text(encoding="utf-8")
    target.write_text("{bad", encoding="utf-8")
    with pytest.raises(CheckpointInvalid):
        store.load(result["state"].workflow_id)
    payload = json.loads(raw)
    payload["checkpoint_schema_version"] = "future-version"
    target.write_text(json.dumps(payload), encoding="utf-8")
    with pytest.raises(CheckpointInvalid, match="version"):
        store.load(result["state"].workflow_id)
    payload = json.loads(raw)
    payload["workflow_status"] = "RUNNING"
    target.write_text(json.dumps(payload), encoding="utf-8")
    changed = Document(source)
    changed.add_paragraph("template changed")
    changed.save(source)
    with pytest.raises(CheckpointInvalid, match="hash"):
        DocumentWorkflow(DemoRAGClient()).resume(result["state"].workflow_id, output=output)


def test_unknown_evidence_blocks_finalization_and_template_overwrite(tmp_path):
    source = tmp_path / "template.docx"
    output = tmp_path / "output"
    two_section_template(source)
    result = DocumentWorkflow(DemoRAGClient()).run(source, output)
    validator = OutputIntegrityValidator()
    result["drafts"][0].evidence_ids.append("E999")
    with pytest.raises(OutputIntegrityError, match="EVIDENCE_MEMBERSHIP_INVALID"):
        validator.validate(source=source, target=output / "other.docx",
                           original_hash=result["trace"]["original_template_hash"],
                           schema=result["schema"], tasks=result["tasks"],
                           drafts=result["drafts"], cache=_restored_cache(output, result),
                           workflow_status="COMPLETE")
    with pytest.raises(OutputIntegrityError, match="ORIGINAL_TEMPLATE_OVERWRITE_BLOCKED"):
        validator.validate(source=source, target=source,
                           original_hash=result["trace"]["original_template_hash"],
                           schema=result["schema"], tasks=result["tasks"],
                           drafts=result["drafts"], cache=_restored_cache(output, result),
                           workflow_status="COMPLETE")


def test_required_section_and_human_review_guard(tmp_path):
    source = tmp_path / "template.docx"
    output = tmp_path / "output"
    two_section_template(source)
    result = DocumentWorkflow(DemoRAGClient()).run(source, output)
    validator = OutputIntegrityValidator()
    arguments = dict(source=source, target=output / "other.docx",
                     original_hash=result["trace"]["original_template_hash"],
                     schema=result["schema"], tasks=result["tasks"],
                     cache=_restored_cache(output, result), workflow_status="COMPLETE")
    with pytest.raises(OutputIntegrityError, match="REQUIRED_SECTION_INVALID"):
        validator.validate(**arguments, drafts=result["drafts"][:1])
    result["drafts"][0].requires_human_review = False
    with pytest.raises(OutputIntegrityError, match="HUMAN_REVIEW_REQUIRED"):
        validator.validate(**arguments, drafts=result["drafts"])
    with pytest.raises(FileExistsError):
        DocumentWorkflow(DemoRAGClient()).run(source, output)


def test_checkpoint_replace_failure_keeps_previous_valid_state(monkeypatch, tmp_path):
    source = tmp_path / "template.docx"
    output = tmp_path / "output"
    two_section_template(source)
    result = DocumentWorkflow(DemoRAGClient()).run(source, output)
    store = CheckpointStore(output / "checkpoints")
    before = store.path(result["state"].workflow_id).read_bytes()

    def refused_replace(*args):
        raise OSError("simulated replacement failure")

    monkeypatch.setattr("app.document_workflow.state.os.replace", refused_replace)
    with pytest.raises(OSError):
        store.save(result["state"])
    assert store.path(result["state"].workflow_id).read_bytes() == before
    assert not list((output / "checkpoints").glob(".checkpoint.*.tmp"))


def test_workflow_lifecycle_complete_partial_and_failed(monkeypatch, tmp_path):
    from app.document_workflow.integrity import OutputIntegrityValidator
    original_save = CheckpointStore.save
    observed = []

    def record(self, state):
        observed.append(state.workflow_status)
        return original_save(self, state)

    monkeypatch.setattr(CheckpointStore, "save", record)
    complete = tmp_path / "complete.docx"
    two_section_template(complete)
    result = DocumentWorkflow(DemoRAGClient()).run(complete, tmp_path / "complete-output")
    assert result["state"].workflow_status == "COMPLETE"
    assert observed[0] == "PENDING" and "RUNNING" in observed and observed[-1] == "COMPLETE"

    observed.clear()
    partial = tmp_path / "partial.docx"
    document = Document()
    document.add_heading("吞吐能力", level=1)
    document.add_paragraph("{{吞吐能力}}")
    document.save(partial)
    result = DocumentWorkflow(DemoRAGClient()).run(partial, tmp_path / "partial-output")
    assert result["state"].workflow_status == "PARTIAL"
    assert observed[0] == "PENDING" and "RUNNING" in observed and observed[-1] == "PARTIAL"

    observed.clear()
    def invalid(*args, **kwargs):
        raise OutputIntegrityError("EVIDENCE_MEMBERSHIP_INVALID")

    monkeypatch.setattr(OutputIntegrityValidator, "validate", invalid)
    with pytest.raises(OutputIntegrityError):
        DocumentWorkflow(DemoRAGClient()).run(complete, tmp_path / "failed-output")
    assert observed[0] == "PENDING" and "RUNNING" in observed and observed[-1] == "FAILED"


def test_trace_masks_obvious_secret_field_name_but_checkpoint_can_resume(tmp_path):
    source = tmp_path / "template.docx"
    document = Document()
    document.add_heading("Configuration", level=1)
    document.add_paragraph("{{SECRET_TOKEN_ABC}}")
    document.save(source)
    output = tmp_path / "output"
    result = DocumentWorkflow(DemoRAGClient()).run(source, output)
    trace = (output / "execution_trace.json").read_text(encoding="utf-8")
    checkpoint = CheckpointStore(output / "checkpoints").path(result["state"].workflow_id).read_text(encoding="utf-8")
    assert "SECRET_TOKEN_ABC" not in trace
    assert "field_sha256:" in trace
    assert "SECRET_TOKEN_ABC" in checkpoint


def _restored_cache(output, result):
    from app.document_workflow.evidence import EvidenceCache
    cache = EvidenceCache(output, output / "evidence_store.json")
    cache.restore(result["state"].evidence_snapshot)
    return cache


def test_http_retry_is_limited_and_schema_failure_is_not_retried(monkeypatch):
    calls = []
    content = "项目背景：安全示例"
    hit = {"chunk_id": "c1", "document_id": "d1", "section_id": "s1",
           "section_path": ["项目背景"], "page_number": 1, "content": content,
           "content_hash": hashlib.sha256(content.encode()).hexdigest(),
           "similarity": 0.8, "rank": 1}

    def server_then_ok(url, *, json, timeout):
        calls.append(url)
        if len(calls) == 1:
            return httpx.Response(503, request=httpx.Request("POST", url))
        return httpx.Response(200, json={"query": json["query"], "results": [hit]},
                              request=httpx.Request("POST", url))

    monkeypatch.setattr(httpx, "post", server_then_ok)
    client = HTTPRetrieveClient("http://localhost:8000", retry_limit=1)
    assert len(client.retrieve("项目背景", 5)["results"]) == 1
    assert len(calls) == 2

    def invalid(url, *, json, timeout):
        calls.append(url)
        return httpx.Response(200, json={"query": json["query"], "results": [{}]},
                              request=httpx.Request("POST", url))

    monkeypatch.setattr(httpx, "post", invalid)
    with pytest.raises(ValueError, match="malformed"):
        client.retrieve("项目背景", 5)
    assert len(calls) == 3
