from __future__ import annotations

import base64
import hashlib
import json
from datetime import datetime, timezone
from pathlib import Path

from docx import Document

from app.change_impact_review import (
    ChangeImpactReviewFacade,
    ChangeTaskStatus,
    PatchCandidate,
    PatchOperation,
    PatchReviewAction,
)
from app.document_workflow.evidence_models import Evidence
from app.document_workflow.rag import HTTPRetrieveClient


class Response:
    def __init__(self, payload):
        self.payload = payload
        self.status_code = 200

    def raise_for_status(self):
        return None

    def json(self):
        return self.payload


def test_existing_rag_http_client_builds_and_activates_candidate(monkeypatch) -> None:
    calls = []

    def post(url, *, json, timeout):
        calls.append((url, json))
        if url.endswith("/engineering/candidates/build"):
            assert base64.b64decode(json["source_base64"], validate=True) == b"docx-bytes"
            return Response({"candidate_id": "candidate-1", "status": "VALIDATED"})
        if url.endswith("/activate"):
            return Response({"candidate_id": "candidate-1", "status": "ACTIVE"})
        if url.endswith("/engineering/versions/retrieve"):
            return Response({"results": [{"version_id": "design-v2"}]})
        raise AssertionError(url)

    monkeypatch.setattr("app.document_workflow.rag.httpx.post", post)
    client = HTTPRetrieveClient("http://rag.local", timeout=5)
    built = client.build_candidate_version(
        document_id="system_design",
        project_id="PAYMENT",
        document_type="DESIGN",
        title="系统设计说明书",
        version_id="design-v2",
        version_label="V2.0",
        source_bytes=b"docx-bytes",
        source_name="candidate.docx",
        profile={"organization_id": "demo_company_a"},
        expected_contents=["1000 并发"],
    )
    assert built["status"] == "VALIDATED"
    assert client.activate_candidate_version("candidate-1")["status"] == "ACTIVE"
    assert client.retrieve_candidate_versions([1.0, 0.0], {"active_only": True})["results"][0]["version_id"] == "design-v2"


class VersionClient:
    def __init__(self, fail: bool = False):
        self.fail = fail
        self.build_calls = 0
        self.activate_calls = 0

    def build_candidate_version(self, **payload):
        self.build_calls += 1
        return {
            "candidate_id": "candidate-1",
            "status": "FAILED" if self.fail else "VALIDATED",
            "failure_reason": "CANDIDATE_BUILD_FAILED" if self.fail else None,
            "validations": {"docx": not self.fail},
        }

    def activate_candidate_version(self, candidate_id):
        self.activate_calls += 1
        return {"candidate_id": candidate_id, "status": "ACTIVE", "validations": {"safe_activation": True}}


def prepared_task(tmp_path: Path, version_client: VersionClient):
    original = "DES-014 面向 REQ-023：连接池支持 500 并发，并同步返回结果。"
    document = Document()
    document.add_heading("系统设计", level=1)
    document.add_paragraph(original)
    source = tmp_path / "source.docx"
    document.save(source)
    ev = Evidence(
        title="需求规格说明书",
        content="REQ-023 并发提升到 1000",
        organization="星海软件科技有限公司（虚构）",
        source_type="RAG",
        document_id="requirements",
        project_id="PAYMENT",
        document_type="REQUIREMENT",
        version_id="requirements-v2",
        version_label="V2.0",
        version_status="ACTIVE",
        chunk_id="req-v2:23",
        section="REQUIREMENTS",
        section_path=["REQUIREMENTS"],
        page_number=1,
        retrieved_at=datetime.now(timezone.utc),
    )
    scope = {"project_ids": ["PAYMENT"], "active_only": True}
    fingerprint = hashlib.sha256(
        json.dumps(scope, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode()
    ).hexdigest()
    patch = PatchCandidate.create(
        target_document_id="system_design",
        base_version_id="design-v1",
        target_section_id="DESIGN",
        target_anchor="paragraph:1",
        operation=PatchOperation.REPLACE_PARAGRAPH,
        original_content=original,
        proposed_content="DES-014 面向 REQ-023：连接池支持 1000 并发，并通过异步任务状态返回结果。",
        reason="需求变更",
        evidence_ids=[ev.evidence_id],
        evidence_content_hashes={ev.evidence_id: ev.content_hash},
        scope_fingerprint=fingerprint,
    )
    facade = ChangeImpactReviewFacade(
        tmp_path / "runtime",
        active_version_for_document=lambda document_id: "requirements-v2",
        version_client=version_client,
    )
    task = facade.create_task(
        organization_id="demo_company_a",
        project_id="PAYMENT",
        scope=scope,
        base_document_id="system_design",
        base_version_id="design-v1",
        source=source,
        impacts=[],
        patches=[patch],
        evidence=[ev],
    )
    facade.review_patch(task.task_id, patch.patch_id, action=PatchReviewAction.APPROVE, reviewer="reviewer")
    facade.apply_reviewed_patches(task.task_id, current_version_id="design-v1")
    return facade, task.task_id, patch.proposed_content


def test_agent_finalizes_candidate_only_after_rag_validation_and_activation(tmp_path: Path) -> None:
    client = VersionClient()
    facade, task_id, expected = prepared_task(tmp_path, client)

    completed = facade.finalize_candidate_version(
        task_id,
        version_id="design-v2",
        version_label="V2.0",
        document_type="DESIGN",
        title="系统设计说明书",
        profile={"organization_id": "demo_company_a"},
    )

    assert completed.status is ChangeTaskStatus.COMPLETED
    assert completed.candidate_version_record["status"] == "ACTIVE"
    assert client.build_calls == 1
    assert client.activate_calls == 1
    assert expected in Document(completed.candidate_path).paragraphs[1].text


def test_agent_does_not_activate_failed_candidate(tmp_path: Path) -> None:
    client = VersionClient(fail=True)
    facade, task_id, _ = prepared_task(tmp_path, client)

    failed = facade.finalize_candidate_version(
        task_id,
        version_id="design-v2",
        version_label="V2.0",
        document_type="DESIGN",
        title="系统设计说明书",
        profile={"organization_id": "demo_company_a"},
    )

    assert failed.status is ChangeTaskStatus.FAILED
    assert failed.failure_reason == "CANDIDATE_BUILD_FAILED"
    assert client.activate_calls == 0

