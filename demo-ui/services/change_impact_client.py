"""UI adapter for data-driven synthetic engineering change cases."""

from __future__ import annotations

import hashlib
import json
import sys
from datetime import datetime, timezone
from typing import Any

from docx import Document

from config import AGENT_ROOT, DemoConfig
from services.demo_cases import DemoCase, load_demo_cases
from services.rag_client import ServiceError


class ChangeImpactClient:
    def __init__(self, config: DemoConfig, demo_case: DemoCase | None = None):
        self.config = config
        self.demo_case = demo_case or load_demo_cases()[config.demo_case_id]
        if str(AGENT_ROOT) not in sys.path:
            sys.path.insert(0, str(AGENT_ROOT))
        from app.change_impact_review import ChangeImpactReviewFacade
        from app.document_workflow.rag import HTTPRetrieveClient

        self.rag = HTTPRetrieveClient(
            config.rag_base_url,
            timeout=config.request_timeout_seconds,
            retry_limit=config.rag_retry_limit,
            session_id=config.session_id,
        )
        self.facade = ChangeImpactReviewFacade(
            config.runtime_root / "change-review",
            active_version_for_document=self._active_version,
            version_client=self.rag,
        )

    def status(self) -> dict[str, Any]:
        return {
            "ready": self.demo_case.inventory_path.is_file(),
            "case_id": self.demo_case.case_id,
            "case_title": self.demo_case.title,
            "description": self.demo_case.description,
            "profile": self.demo_case.organization_id,
            "project": self.demo_case.project_id,
            "classification": "FULLY_SYNTHETIC",
        }

    def _inventory(self) -> dict:
        try:
            payload = json.loads(self.demo_case.inventory_path.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError) as exc:
            raise ServiceError("合成资料清单不可用。", str(exc), "CASE_FIXTURE_MISSING") from exc
        if payload.get("schema_version") != "v4-engineering-inventory-v1":
            raise ServiceError("合成资料清单版本不匹配。", "schema mismatch", "CASE_FIXTURE_INVALID")
        return payload

    def _version_documents(self) -> list[dict]:
        return self.rag.candidate_version_documents()["documents"]

    def _active_version(self, document_id: str) -> str | None:
        for document in self._version_documents():
            if document.get("document_id") == document_id:
                return (document.get("active_version") or {}).get("version_id")
        return None

    def _ensure_seeded(self, inventory: dict) -> None:
        try:
            catalog = self._version_documents()
        except Exception as exc:
            raise ServiceError(
                "版本服务未启用，请配置候选版本存储后重启 RAG。",
                str(exc),
                "VERSION_SERVICE_UNAVAILABLE",
            ) from exc
        known_versions = {
            version["version_id"]
            for document in catalog
            for version in document.get("versions", [])
        }
        for record in inventory["documents"]:
            version_id = record["version_id"]
            if version_id in known_versions:
                continue
            source = self.demo_case.data_root / record["filename"]
            built = self.rag.build_candidate_version(
                document_id=record["document_id"],
                project_id=self.demo_case.project_id,
                document_type=record["document_type"],
                title=record["title"],
                version_id=version_id,
                version_label=record["version_label"],
                source_bytes=source.read_bytes(),
                source_name=record["filename"],
                profile=inventory["profile"],
                expected_contents=[],
            )
            if built.get("status") == "FAILED":
                raise ServiceError(
                    "合成文档版本构建失败，旧有效版本保持不变。",
                    str(built.get("failure_reason")),
                    "SEED_BUILD_FAILED",
                )
            if built.get("status") == "VALIDATED":
                activated = self.rag.activate_candidate_version(str(built["candidate_id"]))
                if activated.get("status") != "ACTIVE":
                    raise ServiceError(
                        "合成文档版本激活失败，旧有效版本保持不变。",
                        str(activated.get("failure_reason")),
                        "SEED_ACTIVATION_FAILED",
                    )
            known_versions.add(version_id)

    def ensure_documents(self) -> None:
        """Seed only this session's existing synthetic case documents."""
        self._ensure_seeded(self._inventory())

    @staticmethod
    def _scope_fingerprint(scope: dict) -> str:
        canonical = json.dumps(scope, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
        return hashlib.sha256(canonical.encode("utf-8")).hexdigest()

    def prepare_demo(self) -> dict[str, Any]:
        return self._prepare()

    def prepare_custom(self, requirement_id: str, proposed_content: str) -> dict[str, Any]:
        """Review a session-only revision of the current case's traced requirement."""
        if requirement_id != self.demo_case.changed_external_identifier:
            raise ValueError("当前需求没有可安全复用的设计与测试追踪关系")
        proposed_content = proposed_content.strip()
        if not proposed_content or len(proposed_content) > 4000 or requirement_id not in proposed_content:
            raise ValueError("请输入包含所选工程编号的需求内容（不超过 4000 字）")
        return self._prepare(custom_content=proposed_content)

    def _prepare(self, *, custom_content: str | None = None) -> dict[str, Any]:
        from app.change_impact_review import PatchCandidate, PatchOperation
        from app.document_workflow.configuration import DocumentWorkflowConfig
        from app.document_workflow.evidence_models import Evidence, normalize_text, sha256_text
        from app.document_workflow.evidence_selection import EvidenceSelectionPolicy, EvidenceSelector
        from app.document_workflow.scope import WorkflowScope

        case = self.demo_case
        inventory = self._inventory()
        self._ensure_seeded(inventory)
        items = inventory["items"]
        current_items = [item for item in items if item["version_id"] == case.requirement_new_version_id]
        baseline = next(
            item for item in current_items
            if item["external_identifier"] == case.changed_external_identifier
        )
        if custom_content is None:
            old_items = [item for item in items if item["version_id"] == case.requirement_old_version_id]
            new_items = current_items
            changed = baseline
            evidence_item = changed
            active_items = [
                item for item in items
                if item["version_id"] != case.requirement_old_version_id
            ]
            trace_links = inventory["trace_links"]
        else:
            if normalize_text(custom_content) == normalize_text(baseline["content"]):
                raise ValueError("修改后内容与当前需求相同")
            identity = hashlib.sha256(
                f"{self.config.session_id}|{baseline['item_id']}|{custom_content}".encode("utf-8")
            ).hexdigest()[:20]
            changed = {
                **baseline,
                "item_id": f"item_{identity}",
                "version_id": f"session-{identity}",
                "content": custom_content,
                "content_hash": sha256_text(normalize_text(custom_content)),
                "metadata": {**baseline["metadata"], "origin": "SESSION_INPUT"},
            }
            old_items = current_items
            new_items = [
                changed if item["item_id"] == baseline["item_id"] else item
                for item in current_items
            ]
            evidence_item = baseline
            active_items = [
                item for item in items
                if item["version_id"] != case.requirement_old_version_id
                and item["item_id"] != baseline["item_id"]
            ] + [changed]
            trace_links = [
                {**link, "source_item_id": changed["item_id"]}
                if link["source_item_id"] == baseline["item_id"] else link
                for link in inventory["trace_links"]
            ]
        changes = self.rag.diff_engineering_items(old_items, new_items)["changes"]
        scope = {"project_ids": [case.project_id], "active_only": True}
        dense = self.rag.search_candidate_versions(changed["content"], scope, top_k=20)["results"]
        by_location = {
            (item["document_id"], item["version_id"], item["section_id"]): item
            for item in active_items
        }
        dense_ids: list[str] = []
        evidence_by_item: dict[str, list[str]] = {}
        evidence_objects: dict[str, Evidence] = {}
        for row in dense:
            item = by_location.get((row["document_id"], row["version_id"], row["section_id"]))
            if item is None or item["item_id"] == changed["item_id"] or item["document_id"] == changed["document_id"]:
                continue
            if item["item_id"] not in dense_ids:
                dense_ids.append(item["item_id"])
            evidence = Evidence(
                title=row.get("document_title") or row["document_id"],
                content=row["text"],
                organization=case.organization_name,
                source_type="RAG",
                document_id=row["document_id"],
                project_id=row["project_id"],
                document_type=row["document_type"],
                version_id=row["version_id"],
                version_label=row["version_label"],
                version_status=row["version_status"],
                chunk_id=row["chunk_id"],
                query=changed["content"],
                section=row["section_id"],
                section_path=row["section_path"],
                page_number=row["page_number"],
                retrieved_at=datetime.now(timezone.utc),
            )
            assert evidence.evidence_id
            evidence_objects[evidence.evidence_id] = evidence
            evidence_by_item.setdefault(item["item_id"], []).append(evidence.evidence_id)
        impacts = self.rag.discover_engineering_impacts(
            changed["item_id"],
            active_items,
            trace_links,
            dense_item_ids=dense_ids,
            evidence_by_item=evidence_by_item,
        )["impacts"]
        version_record = next(
            record for record in inventory["documents"]
            if record["version_id"] == case.requirement_new_version_id
        )
        requirement_evidence = Evidence(
            title=version_record["title"],
            content=evidence_item["content"],
            organization=case.organization_name,
            source_type="RAG",
            document_id=evidence_item["document_id"],
            project_id=evidence_item["project_id"],
            document_type=evidence_item["item_type"],
            version_id=evidence_item["version_id"],
            version_label=version_record["version_label"],
            version_status="ACTIVE",
            chunk_id=f"{evidence_item['version_id']}:{evidence_item['section_id']}:{evidence_item['item_id']}",
            query=changed["external_identifier"],
            section=evidence_item["section_id"],
            section_path=evidence_item["metadata"]["section_path"],
            page_number=1,
            retrieved_at=datetime.now(timezone.utc),
        )
        assert requirement_evidence.evidence_id and requirement_evidence.content_hash
        evidence_objects[requirement_evidence.evidence_id] = requirement_evidence
        maximum = DocumentWorkflowConfig.from_env().max_evidence_count
        selection = EvidenceSelector(EvidenceSelectionPolicy(maximum)).select(
            list(evidence_objects.values()),
            WorkflowScope.from_dict(scope),
            required_evidence_ids=(requirement_evidence.evidence_id,),
        )
        selected_evidence_ids = {
            item.evidence_id for item in selection.selected if item.evidence_id
        }
        if requirement_evidence.evidence_id not in selected_evidence_ids:
            raise RuntimeError("critical requirement Evidence was not selected")

        patch_source = case.data_root / case.patch_document_filename
        patch_document = Document(patch_source)
        paragraph_index = next(
            index for index, paragraph in enumerate(patch_document.paragraphs)
            if case.patch_target_external_identifier in paragraph.text
        )
        original = patch_document.paragraphs[paragraph_index].text
        patch = PatchCandidate.create(
            target_document_id=case.patch_target_document_id,
            base_version_id=case.patch_base_version_id,
            target_section_id=case.patch_target_section_id,
            target_anchor=f"paragraph:{paragraph_index}",
            operation=PatchOperation.REPLACE_PARAGRAPH,
            original_content=original,
            proposed_content=(
                case.patch_proposed_content if custom_content is None
                else original.rstrip() + "\n\n待审核的需求变更（用户输入）： " + custom_content
            ),
            reason=(
                case.patch_reason if custom_content is None
                else f"根据用户本次提出的 {case.changed_external_identifier} 变更形成的待审核建议；引用依据为当前有效资料"
            ),
            evidence_ids=[requirement_evidence.evidence_id],
            evidence_content_hashes={
                requirement_evidence.evidence_id: requirement_evidence.content_hash
            },
            scope_fingerprint=self._scope_fingerprint(scope),
        )
        state = self.facade.create_task(
            organization_id=case.organization_id,
            project_id=case.project_id,
            scope=scope,
            base_document_id=case.patch_target_document_id,
            base_version_id=case.patch_base_version_id,
            source=patch_source,
            impacts=impacts,
            patches=[patch],
            evidence=list(evidence_objects.values()),
            rag_calls=3,
        )
        result = {
            "case_id": case.case_id,
            "state": state.model_dump(mode="json"),
            "changes": changes,
            "items": {item["item_id"]: item for item in active_items},
            "evidence": {
                evidence_id: item.model_dump(mode="json")
                for evidence_id, item in evidence_objects.items()
            },
            "evidence_selection": {
                **selection.stats.model_dump(),
                "max_evidence_count": maximum,
                "selected_evidence_ids": sorted(selected_evidence_ids),
            },
        }
        if custom_content is not None:
            result["custom_change"] = {
                "source": "用户输入",
                "requirement_id": case.changed_external_identifier,
                "old_content": baseline["content"],
                "new_content": custom_content,
                "publication_available": False,
            }
        return result

    def review_patch(
        self,
        task_id: str,
        patch_id: str,
        *,
        action: str,
        reviewer: str,
        comment: str = "",
        edited_content: str | None = None,
    ) -> dict:
        from app.change_impact_review import PatchReviewAction

        state = self.facade.review_patch(
            task_id,
            patch_id,
            action=PatchReviewAction(action),
            reviewer=reviewer,
            comment=comment,
            edited_content=edited_content,
        )
        return state.model_dump(mode="json")

    def apply(self, task_id: str) -> dict:
        current_version_id = self._active_version(self.demo_case.patch_target_document_id) or ""
        state, results = self.facade.apply_reviewed_patches(
            task_id, current_version_id=current_version_id
        )
        return {
            "state": state.model_dump(mode="json"),
            "apply_results": [item.model_dump(mode="json") for item in results],
        }

    def publish(self, task_id: str) -> dict:
        inventory = self._inventory()
        document_type = "系统设计说明书"
        for row in self._version_documents():
            if row.get("document_id") == self.demo_case.patch_target_document_id:
                document_type = str(row.get("document_type") or document_type)
                break
        state = self.facade.finalize_candidate_version(
            task_id,
            version_id=self.demo_case.candidate_version_id,
            version_label=self.demo_case.candidate_version_label,
            document_type=document_type,
            title=self.demo_case.candidate_title,
            profile=inventory["profile"],
        )
        return state.model_dump(mode="json")
