from __future__ import annotations

import base64
from pathlib import Path

import numpy as np
from docx import Document
from fastapi.testclient import TestClient

from src.document_lifecycle import RetrievalScope, VersionStatus
from src.engineering_change import CandidateStatus, CandidateVersionService, OrganizationProfile
from src.rd_v2_api import create_app
from tests.test_rd_v3_runtime_api import _runtime


class FixedEmbedder:
    def encode(self, texts):
        return np.asarray([[1.0, 0.0] for _ in texts], dtype=np.float32)


class FailingEmbedder:
    def encode(self, texts):
        raise RuntimeError("synthetic embedding failure")


def profile() -> OrganizationProfile:
    return OrganizationProfile(
        organization_id="demo_company_a",
        identifier_patterns={"DESIGN": [r"DES-\d{3}"]},
        document_type_mapping={"系统设计说明书": "DESIGN"},
        section_aliases={"系统设计": "DESIGN"},
        version_status_mapping={"有效": "ACTIVE", "历史": "SUPERSEDED"},
    )


def document_bytes(path: Path, body: str) -> bytes:
    document = Document()
    document.add_heading("系统设计", level=1)
    document.add_paragraph(body)
    document.save(path)
    return path.read_bytes()


def build(
    service: CandidateVersionService,
    source: bytes,
    *,
    version_id: str,
    label: str,
    expected: str,
):
    return service.build_candidate(
        document_id="system_design",
        project_id="PAYMENT",
        document_type="DESIGN",
        title="系统设计说明书",
        version_id=version_id,
        version_label=label,
        source_bytes=source,
        source_name=f"{version_id}.docx",
        profile=profile(),
        expected_contents=[expected],
    )


def test_candidate_build_then_safe_activation_preserves_history(tmp_path: Path) -> None:
    v1_text = "DES-014 支持 500 并发并同步返回结果。"
    v2_text = "DES-014 支持 1000 并发，并通过异步任务状态返回结果。"
    service = CandidateVersionService(tmp_path / "store", FixedEmbedder())

    initial = build(service, document_bytes(tmp_path / "v1.docx", v1_text), version_id="design-v1", label="V1.0", expected=v1_text)
    assert initial.status is CandidateStatus.VALIDATED
    active_v1 = service.activate(initial.candidate_id)
    assert active_v1.status is CandidateStatus.ACTIVE
    assert service.lifecycle.catalog.active_version("system_design").version_id == "design-v1"

    candidate = build(service, document_bytes(tmp_path / "v2.docx", v2_text), version_id="design-v2", label="V2.0", expected=v2_text)
    assert candidate.status is CandidateStatus.VALIDATED
    assert all(candidate.validations.values())
    assert service.lifecycle.catalog.active_version("system_design").version_id == "design-v1"

    activated = service.activate(candidate.candidate_id)
    assert activated.status is CandidateStatus.ACTIVE
    assert service.lifecycle.catalog.active_version("system_design").version_id == "design-v2"
    assert service.lifecycle.catalog.by_version["design-v1"].status is VersionStatus.SUPERSEDED

    current = service.retrieve([1.0, 0.0], RetrievalScope(document_ids=["system_design"]), top_k=5)
    historical = service.retrieve(
        [1.0, 0.0],
        RetrievalScope(version_ids=["design-v1"], active_only=False),
        top_k=5,
    )
    assert {row["version_id"] for row in current} == {"design-v2"}
    assert {row["version_id"] for row in historical} == {"design-v1"}


def test_failed_candidate_never_replaces_current_active_version(tmp_path: Path) -> None:
    v1_text = "DES-014 支持 500 并发并同步返回结果。"
    service = CandidateVersionService(tmp_path / "store", FixedEmbedder())
    initial = build(service, document_bytes(tmp_path / "v1.docx", v1_text), version_id="design-v1", label="V1.0", expected=v1_text)
    service.activate(initial.candidate_id)

    failing = CandidateVersionService(tmp_path / "store", FailingEmbedder())
    failed = build(
        failing,
        document_bytes(tmp_path / "v2.docx", "DES-014 损坏构建 1000 并发。"),
        version_id="design-v2",
        label="V2.0",
        expected="DES-014 损坏构建 1000 并发。",
    )

    assert failed.status is CandidateStatus.FAILED
    assert failed.failure_reason == "CANDIDATE_BUILD_FAILED"
    assert failing.lifecycle.catalog.active_version("system_design").version_id == "design-v1"
    assert "design-v2" not in failing.lifecycle.catalog.by_version


def test_candidate_version_http_build_activate_and_retrieve(tmp_path: Path) -> None:
    text = "DES-014 支持 1000 并发和异步任务状态。"
    source = document_bytes(tmp_path / "v1.docx", text)
    versions = CandidateVersionService(tmp_path / "store", FixedEmbedder())

    with TestClient(create_app(_runtime(tmp_path / "runtime"), candidate_service=versions)) as client:
        built = client.post("/engineering/candidates/build", json={
            "document_id": "system_design",
            "project_id": "PAYMENT",
            "document_type": "DESIGN",
            "title": "系统设计说明书",
            "version_id": "design-v1",
            "version_label": "V1.0",
            "source_base64": base64.b64encode(source).decode("ascii"),
            "source_name": "design-v1.docx",
            "profile": profile().model_dump(mode="json"),
            "expected_contents": [text],
        })
        assert built.status_code == 200
        candidate_id = built.json()["candidate_id"]
        assert built.json()["status"] == "VALIDATED"

        activated = client.post(f"/engineering/candidates/{candidate_id}/activate")
        assert activated.status_code == 200
        assert activated.json()["status"] == "ACTIVE"

        retrieved = client.post("/engineering/versions/retrieve", json={
            "query_vector": [1.0, 0.0],
            "scope": {"document_ids": ["system_design"], "active_only": True},
            "top_k": 5,
        })
        assert retrieved.status_code == 200
        assert {row["version_id"] for row in retrieved.json()["results"]} == {"design-v1"}

