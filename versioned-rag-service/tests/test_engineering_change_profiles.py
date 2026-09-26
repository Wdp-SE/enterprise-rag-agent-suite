from pathlib import Path

import pytest
from docx import Document

from src.engineering_change import (
    DocxEngineeringParser,
    EngineeringItemFactory,
    EngineeringItemType,
    IdentifierExtractor,
    OrganizationProfile,
)


def company_a_profile() -> OrganizationProfile:
    return OrganizationProfile(
        organization_id="demo_company_a",
        identifier_patterns={
            "REQUIREMENT": [r"REQ-\d{3}"],
            "DESIGN": [r"DES-\d{3}"],
            "API": [r"API-\d{3}"],
            "TEST_CASE": [r"TC-\d{3}"],
            "RUNBOOK": [r"OPS-\d{3}"],
        },
        document_type_mapping={
            "需求规格说明书": "REQUIREMENT",
            "系统设计说明书": "DESIGN",
            "接口规范": "API",
            "测试用例": "TEST_CASE",
            "运维手册": "RUNBOOK",
        },
        section_aliases={"功能需求": "REQUIREMENTS", "验收测试": "ACCEPTANCE"},
        version_status_mapping={"有效": "ACTIVE", "历史": "SUPERSEDED"},
    )


def company_b_profile() -> OrganizationProfile:
    return OrganizationProfile(
        organization_id="demo_company_b",
        identifier_patterns={
            "REQUIREMENT": [r"PRD-[A-Z]+-\d{3}"],
            "TEST_CASE": [r"CASE-[A-Z]+-\d{3}", r"TEST\.[A-Z]+\.\d{3}"],
        },
        document_type_mapping={"产品需求": "REQUIREMENT", "验收方案": "TEST_CASE"},
        section_aliases={"业务需求": "REQUIREMENTS", "验收场景": "ACCEPTANCE"},
        version_status_mapping={"当前": "ACTIVE", "归档": "SUPERSEDED"},
    )


def test_profiles_change_enterprise_syntax_without_changing_extractor() -> None:
    extractor = IdentifierExtractor()

    company_a = extractor.extract("REQ-023 由 TC-102 验证", company_a_profile())
    company_b = extractor.extract(
        "PRD-PAY-023 由 CASE-PAY-008 和 TEST.AUTH.018 验证", company_b_profile()
    )

    assert [(item.external_identifier, item.item_type) for item in company_a] == [
        ("REQ-023", EngineeringItemType.REQUIREMENT),
        ("TC-102", EngineeringItemType.TEST_CASE),
    ]
    assert [(item.external_identifier, item.item_type) for item in company_b] == [
        ("PRD-PAY-023", EngineeringItemType.REQUIREMENT),
        ("CASE-PAY-008", EngineeringItemType.TEST_CASE),
        ("TEST.AUTH.018", EngineeringItemType.TEST_CASE),
    ]


def test_profile_maps_document_section_and_version_status() -> None:
    a = company_a_profile()
    b = company_b_profile()

    assert a.document_type("需求规格说明书") is EngineeringItemType.REQUIREMENT
    assert b.document_type("产品需求") is EngineeringItemType.REQUIREMENT
    assert a.canonical_section("功能需求") == "REQUIREMENTS"
    assert b.canonical_section("业务需求") == "REQUIREMENTS"
    assert a.version_status("有效") == "ACTIVE"
    assert b.version_status("归档") == "SUPERSEDED"


def test_invalid_identifier_pattern_fails_profile_validation() -> None:
    with pytest.raises(ValueError, match="identifier pattern"):
        OrganizationProfile(
            organization_id="broken",
            identifier_patterns={"REQUIREMENT": ["["]},
            document_type_mapping={"需求": "REQUIREMENT"},
            section_aliases={},
            version_status_mapping={"有效": "ACTIVE"},
        )


def test_docx_parser_and_item_factory_cover_heading_paragraph_and_table(tmp_path: Path) -> None:
    source = tmp_path / "requirements_v2.docx"
    document = Document()
    document.add_heading("项目需求", level=1)
    document.add_heading("功能需求", level=2)
    document.add_paragraph("REQ-023 最大并发提升到 1000，并新增吞吐量 200 MB/s。")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "关联测试"
    table.cell(0, 1).text = "TC-102"
    table.cell(1, 0).text = "说明"
    table.cell(1, 1).text = "持续 30 分钟"
    document.save(source)

    sections = DocxEngineeringParser().parse(source, company_a_profile())
    items = EngineeringItemFactory(IdentifierExtractor()).build(
        sections=sections,
        profile=company_a_profile(),
        project_id="PAYMENT",
        document_id="requirements",
        version_id="requirements@2",
        document_type="需求规格说明书",
    )

    assert sections[0].section_path == ["项目需求", "REQUIREMENTS"]
    assert "关联测试 | TC-102" in sections[0].content
    assert [item.external_identifier for item in items] == ["REQ-023", "TC-102"]
    assert items[0].item_type is EngineeringItemType.REQUIREMENT
    assert items[0].organization_id == "demo_company_a"
    assert items[0].section_id == sections[0].section_id
    assert items[0].item_id == EngineeringItemFactory(IdentifierExtractor()).build(
        sections=sections,
        profile=company_a_profile(),
        project_id="PAYMENT",
        document_id="requirements",
        version_id="requirements@2",
        document_type="需求规格说明书",
    )[0].item_id

