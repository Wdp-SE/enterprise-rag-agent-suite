from __future__ import annotations

from docx import Document
from fastapi.testclient import TestClient

from src.engineering_change import CandidateVersionService
from src.rd_v2_api import create_app
from tests.test_candidate_version_service import FixedEmbedder, profile
from tests.test_rd_v3_runtime_api import _runtime


class AnswerGenerator:
    def __init__(self, page_number: int):
        self.page_number = page_number
        self.context = ""

    def generate(self, *, question: str, context: str) -> dict:
        self.context = context
        return {
            "final_answer": "当前最大并发为 1000。",
            "relevant_sources": [
                {"document_id": "payment_requirement", "page_number": self.page_number},
                {"document_id": "audit_runbook", "page_number": 1},
            ],
        }


def _add_document(service, tmp_path, document_id: str, project_id: str, body: str) -> None:
    source = tmp_path / f"{document_id}.docx"
    document = Document()
    document.add_heading("系统设计", level=1)
    document.add_paragraph(body)
    document.save(source)
    record = service.build_candidate(
        document_id=document_id,
        project_id=project_id,
        document_type="系统设计说明书",
        title=document_id,
        version_id=f"{document_id}-v1",
        version_label="V1.0",
        source_bytes=source.read_bytes(),
        source_name=source.name,
        profile=profile(),
        expected_contents=[body],
    )
    assert record.status == "VALIDATED"
    assert service.activate(record.candidate_id).status == "ACTIVE"


def test_candidate_query_uses_session_scope_real_citations_and_budget(monkeypatch, tmp_path):
    monkeypatch.setenv("APP_ENV", "public_demo")
    monkeypatch.setenv("MAX_LLM_CALLS_PER_SESSION", "1")
    versions = CandidateVersionService(tmp_path / "versions", FixedEmbedder())
    _add_document(versions, tmp_path, "payment_requirement", "PAYMENT", "DES-014 最大并发为 1000。")
    _add_document(versions, tmp_path, "audit_runbook", "AUDIT", "DES-031 日志保留 30 天。")
    hit = versions.search_text("最大并发量是多少", top_k=1, scope=None)[0]
    generator = AnswerGenerator(hit["page_number"])
    runtime = _runtime(tmp_path / "runtime")
    runtime.generator = generator

    request = {"question": "最大并发量是多少？", "scope": {"project_ids": ["PAYMENT"], "active_only": True}}
    with TestClient(create_app(runtime, candidate_service=versions)) as client:
        assert client.post("/engineering/versions/query", json=request).status_code == 400
        headers = {"X-Demo-Session-ID": "session_12345678"}
        response = client.post("/engineering/versions/query", json=request, headers=headers)
        assert response.status_code == 200
        body = response.json()
        assert body["status"] == "OK"
        assert body["answer"] == "当前最大并发为 1000。"
        assert [source["document_id"] for source in body["sources"]] == ["payment_requirement"]
        assert "最大并发为 1000" in body["sources"][0]["content"]
        assert "日志保留 30 天" not in generator.context
        assert "document_type:" in generator.context
        assert "version_status: ACTIVE" in generator.context
        assert "不同文档" in generator.context
        assert client.post("/engineering/versions/query", json=request, headers=headers).status_code == 429


def test_candidate_query_without_generator_returns_no_fabricated_answer(monkeypatch, tmp_path):
    monkeypatch.setenv("APP_ENV", "public_demo")
    versions = CandidateVersionService(tmp_path / "versions", FixedEmbedder())
    _add_document(versions, tmp_path, "payment_requirement", "PAYMENT", "REQ-023 最大并发为 1000。")
    runtime = _runtime(tmp_path / "runtime")
    runtime.generator = None
    request = {"question": "最大并发量是多少？", "scope": {"project_ids": ["PAYMENT"], "active_only": True}}
    with TestClient(create_app(runtime, candidate_service=versions)) as client:
        response = client.post(
            "/engineering/versions/query", json=request,
            headers={"X-Demo-Session-ID": "session_12345678"},
        )
    assert response.status_code == 200
    assert response.json()["status"] == "GENERATION_NOT_CONFIGURED"
    assert response.json()["answer"] == "N/A"
    assert response.json()["sources"] == []
