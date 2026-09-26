"""Build the fully synthetic V4 change-impact demo documents and manifest."""

from __future__ import annotations

import json
from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[2]
OUTPUT = ROOT / "project_delivery" / "v4_change_impact_review" / "demo_data"


DOCUMENTS = {
    "requirements_v1.docx": {
        "title": "星海软件支付平台需求规格说明书 V1.0",
        "sections": [
            ("功能需求", [
                "REQ-023 支付批处理最大并发为 500，客户端通过同步接口等待处理结果。",
                "REQ-024 所有批次必须生成审计流水。",
            ]),
        ],
        "table": [("关联测试", "TC-102"), ("关联设计", "DES-014")],
    },
    "requirements_v2.docx": {
        "title": "星海软件支付平台需求规格说明书 V2.0",
        "sections": [
            ("功能需求", [
                "REQ-023 支付批处理最大并发由 500 提升到 1000，吞吐量目标为 200 MB/s，并改为异步任务接口。",
                "REQ-024 所有批次必须生成审计流水。",
                "REQ-025 异步任务完成后必须保留可查询状态。",
            ]),
        ],
        "table": [("关联测试", "TC-102"), ("关联设计", "DES-014"), ("关联接口", "API-008")],
    },
    "system_design_v1.docx": {
        "title": "星海软件支付平台系统设计说明书 V1.0",
        "sections": [
            ("系统设计", [
                "DES-014 面向 REQ-023：批处理服务按最大并发 500 配置连接池，并在同步请求内返回最终结果。",
                "任务状态仅在请求上下文中保存。",
            ]),
        ],
        "table": [("关联需求", "REQ-023"), ("接口", "API-008")],
    },
    "api_spec_v1.docx": {
        "title": "星海软件支付平台接口规范 V1.0",
        "sections": [
            ("接口定义", [
                "API-008 面向 REQ-023：POST /batches 同步返回处理结果，默认超时 30 秒。",
            ]),
        ],
        "table": [("方法", "POST"), ("路径", "/batches"), ("响应", "最终处理结果")],
    },
    "test_cases_v1.docx": {
        "title": "星海软件支付平台测试用例 V1.0",
        "sections": [
            ("验收测试", [
                "TC-102 验证 REQ-023：以 500 并发持续运行 30 分钟，错误率不高于 0.1%。",
            ]),
        ],
        "table": [("前置条件", "测试环境可用"), ("期望", "同步请求全部完成")],
    },
    "runbook_v1.docx": {
        "title": "星海软件支付平台运维手册 V1.0",
        "sections": [
            ("运行手册", [
                "OPS-006 批处理服务容量告警：并发超过 450 时告警，并检查同步接口超时率。",
            ]),
        ],
        "table": [("监控项", "并发数"), ("阈值", "450")],
    },
}


def build_document(path: Path, spec: dict) -> None:
    document = Document()
    document.add_heading(spec["title"], level=1)
    for section, paragraphs in spec["sections"]:
        document.add_heading(section, level=2)
        for paragraph in paragraphs:
            document.add_paragraph(paragraph)
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "字段"
    table.rows[0].cells[1].text = "内容"
    for label, value in spec["table"]:
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = value
    document.save(path)


def main() -> None:
    OUTPUT.mkdir(parents=True, exist_ok=True)
    for filename, spec in DOCUMENTS.items():
        build_document(OUTPUT / filename, spec)
    manifest = {
        "classification": "FULLY_SYNTHETIC",
        "organization_id": "demo_company_a",
        "organization_name": "星海软件科技有限公司（虚构）",
        "project_id": "PAYMENT",
        "documents": list(DOCUMENTS),
        "trace_links": [
            {"source": "REQ-023", "target": "DES-014", "provenance": "EXPLICIT", "status": "CONFIRMED"},
            {"source": "REQ-023", "target": "API-008", "provenance": "EXPLICIT", "status": "CONFIRMED"},
            {"source": "REQ-023", "target": "TC-102", "provenance": "EXPLICIT", "status": "CONFIRMED"},
            {"source": "REQ-023", "target": "OPS-006", "provenance": "SEMANTIC", "status": "SUGGESTED"}
        ]
    }
    (OUTPUT / "manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )


if __name__ == "__main__":
    main()

