"""Build the independent synthetic Case B through the production parser."""

from __future__ import annotations

import json
import sys
from pathlib import Path

from docx import Document

RAG_ROOT = Path(__file__).resolve().parents[1]
WORKSPACE_ROOT = RAG_ROOT.parent
if str(RAG_ROOT) not in sys.path:
    sys.path.insert(0, str(RAG_ROOT))

from src.engineering_change import (
    DocxEngineeringParser,
    EngineeringItemFactory,
    IdentifierExtractor,
    OrganizationProfile,
    TraceLink,
    TraceProvenance,
    TraceStatus,
)


CASE_ROOT = WORKSPACE_ROOT / "project_delivery/public_value_prototype/demo_cases/case_b"
CASE_A_INVENTORY = WORKSPACE_ROOT / "project_delivery/v4_change_impact_review/demo_data/engineering_inventory.json"

DOCUMENTS = {
    "audit_requirements_v1.docx": (
        "audit_requirements", "audit-requirements-v1", "需求规格说明书",
        "远岚审计平台需求规格说明书 V1.0", "数据留存要求",
        ["REQ-071 审计日志在线可检索保存 7 天，超过周期后进入归档。"],
    ),
    "audit_requirements_v2.docx": (
        "audit_requirements", "audit-requirements-v2", "需求规格说明书",
        "远岚审计平台需求规格说明书 V2.0", "合规与留存",
        [
            "REQ-071 审计日志在线可检索保存周期由 7 天调整为 30 天，并按日期分区。",
            "REQ-072 第 31 天起转入低成本归档并保留完整性校验。",
            "关联设计 DES-031，关联测试 TC-207。",
        ],
    ),
    "audit_design_v1.docx": (
        "audit_design", "audit-design-v1", "系统设计说明书",
        "远岚审计平台系统设计说明书 V1.0", "存储与归档设计",
        [
            "DES-031 面向 REQ-071：审计日志按日期分区，在线存储仅保留 7 天，之后转入归档。",
            "索引服务只覆盖在线存储分区。",
        ],
    ),
    "audit_test_cases_v1.docx": (
        "audit_test_cases", "audit-test-v1", "测试用例",
        "远岚审计平台测试用例 V1.0", "合规验收",
        ["TC-207 验证 REQ-071：第 7 天日志仍可查询，第 8 天日志只可从归档读取。"],
    ),
    "audit_runbook_v1.docx": (
        "audit_runbook", "audit-runbook-v1", "运维手册",
        "远岚审计平台运维手册 V1.0", "存储容量巡检",
        ["OPS-031 每日检查在线日志分区容量，并在保存周期接近阈值时告警。"],
    ),
}


def _write_docx(path: Path, title: str, section: str, paragraphs: list[str]) -> None:
    document = Document()
    document.add_heading(title, level=1)
    document.add_heading(section, level=2)
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    document.save(path)


def main() -> None:
    CASE_ROOT.mkdir(parents=True, exist_ok=True)
    base_payload = json.loads(CASE_A_INVENTORY.read_text(encoding="utf-8"))
    profile_payload = dict(base_payload["profile"])
    profile_payload["organization_id"] = "demo_company_b"
    profile = OrganizationProfile.model_validate(profile_payload)
    parser = DocxEngineeringParser()
    factory = EngineeringItemFactory(IdentifierExtractor())
    items = []
    documents = []

    for filename, values in DOCUMENTS.items():
        document_id, version_id, document_type, title, section, paragraphs = values
        source = CASE_ROOT / filename
        _write_docx(source, title, section, paragraphs)
        sections = parser.parse(source, profile)
        extracted = factory.build(
            sections=sections,
            profile=profile,
            project_id="AUDIT",
            document_id=document_id,
            version_id=version_id,
            document_type=document_type,
        )
        items.extend(
            item for item in extracted
            if item.item_type.value == item.metadata["default_document_item_type"]
        )
        documents.append({
            "filename": filename,
            "document_id": document_id,
            "version_id": version_id,
            "version_label": "V2.0" if version_id == "audit-requirements-v2" else "V1.0",
            "document_type": document_type,
            "title": source.stem,
        })

    by_key = {(item.version_id, item.external_identifier): item for item in items}
    requirement = by_key[("audit-requirements-v2", "REQ-071")]
    explicit_targets = [
        by_key[("audit-design-v1", "DES-031")],
        by_key[("audit-test-v1", "TC-207")],
    ]
    links = [
        TraceLink.create(
            source_item_id=requirement.item_id,
            target_item_id=target.item_id,
            provenance=TraceProvenance.EXPLICIT,
            status=TraceStatus.CONFIRMED,
            metadata={"source": "independent synthetic Case B trace"},
        )
        for target in explicit_targets
    ]
    inventory = {
        "schema_version": "v4-engineering-inventory-v1",
        "classification": "FULLY_SYNTHETIC",
        "profile": profile.model_dump(mode="json"),
        "documents": documents,
        "items": [item.model_dump(mode="json") for item in items],
        "trace_links": [link.model_dump(mode="json") for link in links],
        "changed_external_identifier": "REQ-071",
        "patch_target_external_identifier": "DES-031",
    }
    (CASE_ROOT / "engineering_inventory.json").write_text(
        json.dumps(inventory, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    manifest = {
        "classification": "FULLY_SYNTHETIC",
        "organization_id": "demo_company_b",
        "organization_name": "远岚研发中心（虚构）",
        "project_id": "AUDIT",
        "documents": list(DOCUMENTS),
        "change": "日志在线保存周期由 7 天调整为 30 天",
        "confirmed_trace_targets": ["DES-031", "TC-207"],
        "suggested_impact_target": "OPS-031",
    }
    (CASE_ROOT / "manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )


if __name__ == "__main__":
    main()
