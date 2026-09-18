"""Deterministic Agent fresh, resume, and stale-refresh engineering benchmarks."""

from __future__ import annotations

import hashlib
import json
import shutil
import sys
import tempfile
import time
from pathlib import Path

from benchmark_utils import read_json, reduction_percent, stats, utc_now, write_json


SCRIPT_DIR = Path(__file__).resolve().parent
BENCHMARK_ROOT = SCRIPT_DIR.parent
WORKSPACE = BENCHMARK_ROOT.parents[1]
AGENT_ROOT = WORKSPACE / "OpenManus-rag"
RUNTIME_ROOT = BENCHMARK_ROOT / "runtime" / "agent"
FRESH_OUTPUT = BENCHMARK_ROOT / "agent_fresh_run_benchmark.json"
RESUME_OUTPUT = BENCHMARK_ROOT / "agent_resume_benchmark.json"
STALE_OUTPUT = BENCHMARK_ROOT / "agent_stale_refresh_benchmark.json"
WARMUP_RUNS = 5
MEASURED_RUNS = 30

sys.path.insert(0, str(AGENT_ROOT))

from docx import Document  # noqa: E402

from app.document_workflow.configuration import DocumentWorkflowConfig  # noqa: E402
from app.document_workflow.drafting import (  # noqa: E402
    DraftPolicy,
    FieldDraftingService,
)
from app.document_workflow.evidence_models import normalize_text  # noqa: E402
from app.document_workflow.models import DraftingMode, TaskStatus  # noqa: E402
from app.document_workflow.scope import WorkflowScope  # noqa: E402
from app.document_workflow.workflow import DocumentWorkflow  # noqa: E402


class BenchmarkRAGClient:
    """Synthetic retrieval contract with observable calls and controlled versions."""

    def __init__(self, fixture: dict, *, upgraded_documents=(), interrupt_after=None):
        self.fixture = fixture
        self.upgraded_documents = set(upgraded_documents)
        self.interrupt_after = interrupt_after
        self.successful_calls: list[dict] = []
        self.call_elapsed_ms: list[float] = []
        self.attempt_count = 0

    def active_versions(self) -> dict[str, str]:
        return {
            row["document_id"]: self.version_id(row["document_id"])
            for row in self.fixture["sections"]
        }

    def version_id(self, document_id: str) -> str:
        suffix = "2.0" if document_id in self.upgraded_documents else "1.0"
        return f"{document_id}@{suffix}"

    def retrieve(self, query: str, top_k: int, scope: dict | None = None) -> dict:
        self.attempt_count += 1
        if self.interrupt_after is not None and len(self.successful_calls) >= self.interrupt_after:
            raise KeyboardInterrupt("controlled benchmark interruption")
        started = time.perf_counter()
        selected = None
        for row in self.fixture["sections"]:
            if row["title"] in query or row["field"] in query:
                selected = row
                break
        if selected is None:
            raise AssertionError(f"query did not map to a fixture section: {query}")
        expected_scope = {
            "active_only": True,
            "project_ids": [self.fixture["project_id"]],
        }
        if scope != expected_scope:
            raise AssertionError(f"workflow scope changed: {scope}")
        version_id = self.version_id(selected["document_id"])
        version_label = "V2.0" if version_id.endswith("2.0") else "V1.0"
        content = selected["content"]
        if version_id.endswith("2.0"):
            content += " 版本二补充：该章节已按最新有效需求重新确认。"
        content_hash = hashlib.sha256(normalize_text(content).encode("utf-8")).hexdigest()
        response = {
            "query": query,
            "results": [
                {
                    "rank": 1,
                    "similarity": 0.95,
                    "chunk_id": f"{version_id}:{selected['document_id']}",
                    "document_id": selected["document_id"],
                    "project_id": self.fixture["project_id"],
                    "document_type": "REQUIREMENT",
                    "version_id": version_id,
                    "version_label": version_label,
                    "version_status": "ACTIVE",
                    "section_id": selected["title"],
                    "section_path": [selected["title"]],
                    "page_number": int(selected["title"].split()[0]),
                    "content": content,
                    "content_hash": content_hash,
                }
            ][:top_k],
        }
        elapsed_ms = (time.perf_counter() - started) * 1000.0
        self.call_elapsed_ms.append(elapsed_ms)
        self.successful_calls.append(
            {"query": query, "scope": scope, "document_id": selected["document_id"]}
        )
        return response


class TimedDraftingService:
    def __init__(self):
        self.service = FieldDraftingService(
            DraftPolicy(DraftingMode.EXTRACTIVE, "synthetic")
        )
        self.elapsed_ms: list[float] = []

    def draft(self, *args, **kwargs):
        started = time.perf_counter()
        result = self.service.draft(*args, **kwargs)
        self.elapsed_ms.append((time.perf_counter() - started) * 1000.0)
        return result


def build_template(path: Path, fixture: dict) -> Path:
    document = Document()
    for row in fixture["sections"]:
        document.add_heading(row["title"], level=1)
        document.add_paragraph("{{" + row["field"] + "}}")
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(path)
    return path


def workflow_config(root: Path, output: Path) -> DocumentWorkflowConfig:
    return DocumentWorkflowConfig(
        rag_base_url=None,
        rag_top_k=5,
        output_root=root,
        checkpoint_root=output / "checkpoints",
        trace_root=output,
        drafting_mode=DraftingMode.EXTRACTIVE,
        data_classification="synthetic",
    ).validate()


def semantics(result: dict) -> dict:
    return {
        draft.title: [
            {
                "content": field.content,
                "status": field.status.value,
                "evidence_ids": list(field.evidence_ids),
            }
            for field in draft.fields
        ]
        for draft in result["drafts"]
    }


def assert_complete(result: dict, expected_sections: int) -> None:
    if len(result["tasks"]) != expected_sections or len(result["drafts"]) != expected_sections:
        raise AssertionError("workflow did not preserve the required output shape")
    if any(task.status is not TaskStatus.COMPLETE for task in result["tasks"]):
        raise AssertionError("synthetic workflow did not complete every section")
    if len(result["state"].section_drafts) != expected_sections:
        raise AssertionError("checkpoint does not contain every section draft")


def execute_fresh(root: Path, template: Path, fixture: dict, *, upgraded_documents=()) -> tuple[dict, dict]:
    output = root / "output"
    client = BenchmarkRAGClient(fixture, upgraded_documents=upgraded_documents)
    drafting = TimedDraftingService()
    workflow = DocumentWorkflow(
        client,
        config=workflow_config(root, output),
        scope=WorkflowScope(project_ids=(fixture["project_id"],), active_only=True),
        drafting_service=drafting,
    )
    started = time.perf_counter()
    result = workflow.run(template, output)
    elapsed_ms = (time.perf_counter() - started) * 1000.0
    assert_complete(result, len(fixture["sections"]))
    metrics = {
        "elapsed_ms": elapsed_ms,
        "rag_elapsed_ms": sum(client.call_elapsed_ms),
        "drafting_elapsed_ms": sum(drafting.elapsed_ms),
        "rag_calls": len(client.successful_calls),
        "llm_calls": 0,
        "sections_executed": len(result["trace"]["sections"]),
        "fields_executed": sum(len(draft.fields) for draft in result["drafts"]),
        "evidence_count": result["trace"]["total_unique_evidence"],
        "drafted_fields": sum(
            field.status.value == "DRAFTED" for draft in result["drafts"] for field in draft.fields
        ),
        "missing_fields": sum(
            field.status.value != "DRAFTED" for draft in result["drafts"] for field in draft.fields
        ),
    }
    return result, metrics


def prepare_interrupted(root: Path, template: Path, fixture: dict, completed: int) -> tuple[str, Path]:
    output = root / "interrupted"
    client = BenchmarkRAGClient(fixture, interrupt_after=completed)
    workflow = DocumentWorkflow(
        client,
        config=workflow_config(root, output),
        scope=WorkflowScope(project_ids=(fixture["project_id"],), active_only=True),
        drafting_service=TimedDraftingService(),
    )
    try:
        workflow.run(template, output)
    except KeyboardInterrupt:
        pass
    else:
        raise AssertionError("controlled interruption did not occur")
    checkpoints = list((output / "checkpoints").glob("dw_*.json"))
    if len(checkpoints) != 1:
        raise AssertionError("interrupted workflow did not produce one checkpoint")
    payload = json.loads(checkpoints[0].read_text(encoding="utf-8"))
    completed_in_checkpoint = sum(
        row["status"] == "COMPLETE" for row in payload["section_states"].values()
    )
    if completed_in_checkpoint != completed:
        raise AssertionError("checkpoint completion point is incorrect")
    return checkpoints[0].stem, output


def execute_resume(root: Path, template: Path, fixture: dict, workflow_id: str, output: Path) -> tuple[dict, dict]:
    client = BenchmarkRAGClient(fixture)
    drafting = TimedDraftingService()
    workflow = DocumentWorkflow(
        client,
        config=workflow_config(root, output),
        scope=WorkflowScope(project_ids=(fixture["project_id"],), active_only=True),
        drafting_service=drafting,
    )
    started = time.perf_counter()
    result = workflow.resume(workflow_id, output=output, template=template)
    elapsed_ms = (time.perf_counter() - started) * 1000.0
    assert_complete(result, len(fixture["sections"]))
    return result, {
        "elapsed_ms": elapsed_ms,
        "rag_elapsed_ms": sum(client.call_elapsed_ms),
        "drafting_elapsed_ms": sum(drafting.elapsed_ms),
        "rag_calls": len(client.successful_calls),
        "llm_calls": 0,
        "sections_executed": len(result["trace"]["sections"]),
    }


def execute_stale_resume(root: Path, template: Path, fixture: dict, workflow_id: str, output: Path) -> tuple[dict, dict]:
    stale_documents = set(fixture["stale_document_ids"])
    client = BenchmarkRAGClient(fixture, upgraded_documents=stale_documents)
    drafting = TimedDraftingService()
    workflow = DocumentWorkflow(
        client,
        config=workflow_config(root, output),
        scope=WorkflowScope(project_ids=(fixture["project_id"],), active_only=True),
        drafting_service=drafting,
    )
    started = time.perf_counter()
    result = workflow.resume(workflow_id, output=output, template=template)
    elapsed_ms = (time.perf_counter() - started) * 1000.0
    assert_complete(result, len(fixture["sections"]))
    stale_events = [
        row for row in result["trace"]["freshness_events"] if row["freshness"] == "STALE"
    ]
    return result, {
        "elapsed_ms": elapsed_ms,
        "rag_elapsed_ms": sum(client.call_elapsed_ms),
        "drafting_elapsed_ms": sum(drafting.elapsed_ms),
        "rag_calls": len(client.successful_calls),
        "llm_calls": 0,
        "sections_executed": len(result["trace"]["sections"]),
        "stale_evidence_count": len(stale_events),
        "stale_document_ids": sorted({row["document_id"] for row in stale_events}),
    }


def sample_summary(rows: list[dict]) -> dict:
    return {
        "elapsed_ms": stats(row["elapsed_ms"] for row in rows),
        "rag_elapsed_ms": stats(row["rag_elapsed_ms"] for row in rows),
        "drafting_elapsed_ms": stats(row["drafting_elapsed_ms"] for row in rows),
        "rag_calls": sorted({row["rag_calls"] for row in rows}),
        "llm_calls": sorted({row["llm_calls"] for row in rows}),
        "sections_executed": sorted({row["sections_executed"] for row in rows}),
    }


def run_fresh_samples(template: Path, fixture: dict) -> dict:
    for index in range(WARMUP_RUNS):
        with tempfile.TemporaryDirectory(prefix=f"fresh-warm-{index}-", dir=RUNTIME_ROOT) as temp:
            execute_fresh(Path(temp), template, fixture)
    rows = []
    for index in range(MEASURED_RUNS):
        with tempfile.TemporaryDirectory(prefix=f"fresh-{index}-", dir=RUNTIME_ROOT) as temp:
            _, metrics = execute_fresh(Path(temp), template, fixture)
            rows.append(metrics)
    summary = sample_summary(rows)
    representative = rows[0]
    return {
        "schema_version": 1,
        "benchmark": "Agent fresh workflow",
        "captured_at": utc_now(),
        "status": "PASS",
        "benchmark_mode": "OFFLINE_DETERMINISTIC",
        "data_policy": "SYNTHETIC",
        "drafting_mode": "EXTRACTIVE",
        "online_models_called": False,
        "warmup_runs": WARMUP_RUNS,
        "measured_runs": MEASURED_RUNS,
        "counts": {
            "total_sections": len(fixture["sections"]),
            "total_fields": len(fixture["sections"]),
            "total_rag_calls": representative["rag_calls"],
            "llm_calls": 0,
            "evidence_count": representative["evidence_count"],
            "drafted_fields": representative["drafted_fields"],
            "missing_fields": representative["missing_fields"],
        },
        "latency": summary,
        "raw_runs": rows,
        "correctness": "PASS",
    }


def run_resume_samples(template: Path, fixture: dict) -> dict:
    total_sections = len(fixture["sections"])
    scenarios = {}
    for point in fixture["resume_completion_points_percent"]:
        completed = total_sections * int(point) // 100
        if completed <= 0 or completed >= total_sections:
            raise ValueError("resume completion point must be internal to the workflow")

        def one_pair(root: Path, *, reverse: bool) -> tuple[dict, dict]:
            workflow_id, interrupted_output = prepare_interrupted(
                root, template, fixture, completed
            )
            if reverse:
                resumed, resume_metrics = execute_resume(
                    root, template, fixture, workflow_id, interrupted_output
                )
                full, full_metrics = execute_fresh(root / "full", template, fixture)
            else:
                full, full_metrics = execute_fresh(root / "full", template, fixture)
                resumed, resume_metrics = execute_resume(
                    root, template, fixture, workflow_id, interrupted_output
                )
            if semantics(full) != semantics(resumed):
                raise AssertionError("resume output semantics differ from full rerun")
            if full_metrics["rag_calls"] != total_sections:
                raise AssertionError("full rerun did not execute every section")
            if resume_metrics["rag_calls"] != total_sections - completed:
                raise AssertionError("resume did not skip the completed sections")
            return full_metrics, resume_metrics

        for index in range(WARMUP_RUNS):
            with tempfile.TemporaryDirectory(prefix=f"resume-{point}-warm-{index}-", dir=RUNTIME_ROOT) as temp:
                one_pair(Path(temp), reverse=bool(index % 2))
        full_rows = []
        resume_rows = []
        for index in range(MEASURED_RUNS):
            with tempfile.TemporaryDirectory(prefix=f"resume-{point}-{index}-", dir=RUNTIME_ROOT) as temp:
                full_metrics, resume_metrics = one_pair(Path(temp), reverse=bool(index % 2))
                full_rows.append(full_metrics)
                resume_rows.append(resume_metrics)
        full_summary = sample_summary(full_rows)
        resume_summary = sample_summary(resume_rows)
        full_calls = full_rows[0]["rag_calls"]
        resume_calls = resume_rows[0]["rag_calls"]
        scenarios[str(point)] = {
            "completion_point_percent": int(point),
            "sections_completed_before_interruption": completed,
            "sections_skipped_on_resume": completed,
            "full_rerun": full_summary,
            "resume": resume_summary,
            "comparison": {
                "full_rerun_p50_ms": full_summary["elapsed_ms"]["p50"],
                "resume_p50_ms": resume_summary["elapsed_ms"]["p50"],
                "elapsed_reduction_percent": reduction_percent(
                    full_summary["elapsed_ms"]["p50"],
                    resume_summary["elapsed_ms"]["p50"],
                ),
                "full_rerun_rag_calls": full_calls,
                "resume_rag_calls": resume_calls,
                "rag_call_reduction_percent": reduction_percent(full_calls, resume_calls),
                "full_rerun_sections": total_sections,
                "resume_sections": total_sections - completed,
                "section_reexecution_reduction_percent": round(
                    completed / total_sections * 100.0, 3
                ),
            },
            "raw_runs": {"full_rerun": full_rows, "resume": resume_rows},
            "correctness": "PASS",
        }
    return {
        "schema_version": 1,
        "benchmark": "Agent full rerun versus checkpoint resume",
        "captured_at": utc_now(),
        "status": "PASS",
        "benchmark_mode": "OFFLINE_DETERMINISTIC",
        "data_policy": "SYNTHETIC",
        "drafting_mode": "EXTRACTIVE",
        "online_models_called": False,
        "warmup_runs_per_scenario": WARMUP_RUNS,
        "measured_runs_per_scenario": MEASURED_RUNS,
        "total_sections": total_sections,
        "total_fields": total_sections,
        "scenarios": scenarios,
        "correctness": {
            "same_scope": "PASS",
            "same_template": "PASS",
            "same_required_output": "PASS",
            "completed_sections_skipped": "PASS",
        },
    }


def run_stale_samples(template: Path, fixture: dict) -> dict:
    total_sections = len(fixture["sections"])
    stale_count = len(fixture["stale_document_ids"])

    def one_pair(root: Path, *, reverse: bool) -> tuple[dict, dict]:
        seed, _ = execute_fresh(root / "seed", template, fixture)
        workflow_id = seed["state"].workflow_id
        seed_output = Path(seed["state"].output_path)
        if reverse:
            local, local_metrics = execute_stale_resume(
                root / "seed", template, fixture, workflow_id, seed_output
            )
            full, full_metrics = execute_fresh(
                root / "full", template, fixture,
                upgraded_documents=fixture["stale_document_ids"],
            )
        else:
            full, full_metrics = execute_fresh(
                root / "full", template, fixture,
                upgraded_documents=fixture["stale_document_ids"],
            )
            local, local_metrics = execute_stale_resume(
                root / "seed", template, fixture, workflow_id, seed_output
            )
        if semantics(full) != semantics(local):
            raise AssertionError("local refresh output semantics differ from full refresh")
        if full_metrics["rag_calls"] != total_sections or local_metrics["rag_calls"] != stale_count:
            raise AssertionError("refresh call counts do not match affected sections")
        if local_metrics["stale_document_ids"] != sorted(fixture["stale_document_ids"]):
            raise AssertionError("freshness detector selected the wrong documents")
        return full_metrics, local_metrics

    for index in range(WARMUP_RUNS):
        with tempfile.TemporaryDirectory(prefix=f"stale-warm-{index}-", dir=RUNTIME_ROOT) as temp:
            one_pair(Path(temp), reverse=bool(index % 2))
    full_rows = []
    local_rows = []
    for index in range(MEASURED_RUNS):
        with tempfile.TemporaryDirectory(prefix=f"stale-{index}-", dir=RUNTIME_ROOT) as temp:
            full_metrics, local_metrics = one_pair(Path(temp), reverse=bool(index % 2))
            full_rows.append(full_metrics)
            local_rows.append(local_metrics)
    full_summary = sample_summary(full_rows)
    local_summary = sample_summary(local_rows)
    full_calls = full_rows[0]["rag_calls"]
    local_calls = local_rows[0]["rag_calls"]
    preserved = total_sections - stale_count
    return {
        "schema_version": 1,
        "benchmark": "Agent full refresh versus stale local refresh",
        "captured_at": utc_now(),
        "status": "PASS",
        "benchmark_mode": "OFFLINE_DETERMINISTIC",
        "data_policy": "SYNTHETIC",
        "drafting_mode": "EXTRACTIVE",
        "online_models_called": False,
        "warmup_runs": WARMUP_RUNS,
        "measured_runs": MEASURED_RUNS,
        "counts": {
            "total_sections": total_sections,
            "stale_sections": stale_count,
            "refreshed_sections": stale_count,
            "preserved_sections": preserved,
            "refresh_ratio_percent": round(stale_count / total_sections * 100.0, 3),
            "preserved_ratio_percent": round(preserved / total_sections * 100.0, 3),
        },
        "full_refresh": full_summary,
        "local_refresh": local_summary,
        "comparison": {
            "full_refresh_p50_ms": full_summary["elapsed_ms"]["p50"],
            "local_refresh_p50_ms": local_summary["elapsed_ms"]["p50"],
            "elapsed_reduction_percent": reduction_percent(
                full_summary["elapsed_ms"]["p50"],
                local_summary["elapsed_ms"]["p50"],
            ),
            "full_refresh_rag_calls": full_calls,
            "local_refresh_rag_calls": local_calls,
            "rag_call_reduction_percent": reduction_percent(full_calls, local_calls),
            "section_reexecution_reduction_percent": round(
                preserved / total_sections * 100.0, 3
            ),
        },
        "raw_runs": {"full_refresh": full_rows, "local_refresh": local_rows},
        "correctness": {
            "freshness_detected_exact_affected_set": "PASS",
            "preserved_sections_not_reexecuted": "PASS",
            "same_required_output": "PASS",
            "active_version_evidence_after_refresh": "PASS",
        },
    }


def main() -> int:
    fixture = read_json(BENCHMARK_ROOT / "fixtures" / "agent_workflow.json")
    RUNTIME_ROOT.mkdir(parents=True, exist_ok=True)
    template_root = RUNTIME_ROOT / "fixture"
    shutil.rmtree(template_root, ignore_errors=True)
    template = build_template(template_root / "agent_benchmark_template.docx", fixture)
    fresh = run_fresh_samples(template, fixture)
    resume = run_resume_samples(template, fixture)
    stale = run_stale_samples(template, fixture)
    write_json(FRESH_OUTPUT, fresh)
    write_json(RESUME_OUTPUT, resume)
    write_json(STALE_OUTPUT, stale)
    print(
        json.dumps(
            {
                "fresh": fresh["status"],
                "resume": resume["status"],
                "stale_refresh": stale["status"],
            },
            ensure_ascii=False,
            indent=2,
        )
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
