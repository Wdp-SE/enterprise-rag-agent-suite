from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path

import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE

from app.document_workflow.drafting import (
    DraftPolicy,
    FieldDraftingService,
    GeneratedDraft,
)
from app.document_workflow.evidence_models import Evidence, SourceLevel
from app.document_workflow.models import DraftingMode, FieldDraftStatus, FieldTask, TemplateSection
from app.document_workflow.planning import QueryPlanner
from app.document_workflow.scope import WorkflowScope
from app.document_workflow.sufficiency import (
    EvidenceSufficiency,
    EvidenceSufficiencyService,
)
from app.document_workflow.template import TemplateParser, TemplateRenderer


def evidence(content: str, *, chunk: str = "c1") -> Evidence:
    return Evidence(
        title="需求规格说明书",
        content=content,
        organization="研发中心",
        source_type="RAG",
        document_id="REQ-001",
        project_id="P-001",
        document_type="REQUIREMENT",
        version_id="REQ-001@2.0",
        version_label="V2.0",
        version_status="ACTIVE",
        freshness="FRESH",
        chunk_id=chunk,
        section="变更内容",
        section_path=["需求变更", "变更内容"],
        page_number=12,
        source_level=SourceLevel.TIER3,
        retrieved_at=datetime.now(timezone.utc),
    )


def test_scope_is_normalized_fingerprinted_and_rejects_active_history():
    scope = WorkflowScope(project_ids=("P-2", "P-1", "P-1"), active_only=True)
    assert scope.project_ids == ("P-1", "P-2")
    assert WorkflowScope.from_dict(scope.to_dict()).fingerprint() == scope.fingerprint()
    with pytest.raises(ValueError, match="active_only"):
        WorkflowScope(version_ids=("REQ@1",), active_only=True)


def test_query_planner_is_bounded_and_scope_aware():
    section = TemplateSection("s1", "性能变更", 1, 1, None, 0)
    field = FieldTask("f1", "最大并发", "paragraph", True, "s1")
    queries = QueryPlanner().plan(
        section,
        field,
        WorkflowScope(project_ids=("P-001",), document_types=("REQUIREMENT",)),
        [],
    )
    assert 1 <= len(queries) <= 3
    assert queries[0].kind == "PRIMARY"
    assert "性能变更" in queries[0].text and "最大并发" in queries[0].text
    assert any("P-001" in item.text for item in queries)


def test_sufficiency_has_all_three_states():
    field = FieldTask("f1", "变更内容", "paragraph", True, "s1")
    service = EvidenceSufficiencyService()
    assert service.evaluate(field, []).status is EvidenceSufficiency.MISSING
    assert service.evaluate(field, [evidence("短", chunk="short")]).status is EvidenceSufficiency.INSUFFICIENT
    assert service.evaluate(field, [evidence("最大并发由500提升到1000。")]).status is EvidenceSufficiency.SUFFICIENT


def test_extractive_drafting_combines_multiple_evidence():
    field = FieldTask("f1", "变更内容", "paragraph", True, "s1")
    section = TemplateSection("s1", "变更内容", 1, 1, None, 0)
    items = [
        evidence("最大并发由500提升到1000。", chunk="c1"),
        evidence("连接池和限流器需要同步调整。", chunk="c2"),
    ]
    decision = EvidenceSufficiencyService().evaluate(field, items)
    draft = FieldDraftingService().draft(field, section, items, decision)
    assert draft.status is FieldDraftStatus.DRAFTED
    assert len(draft.evidence_ids) == 2
    assert "1000" in draft.content and "连接池" in draft.content


def test_generative_mode_rejects_unapproved_data_and_unknown_citations():
    field = FieldTask("f1", "变更内容", "paragraph", True, "s1")
    section = TemplateSection("s1", "变更内容", 1, 1, None, 0)
    item = evidence("最大并发由500提升到1000。")
    decision = EvidenceSufficiencyService().evaluate(field, [item])

    class Backend:
        def generate(self, **_):
            return GeneratedDraft("生成内容", ("ev_unknown",))

    with pytest.raises(PermissionError):
        FieldDraftingService(DraftPolicy(DraftingMode.GENERATIVE, "internal"), Backend()).draft(
            field, section, [item], decision
        )
    draft = FieldDraftingService(
        DraftPolicy(DraftingMode.GENERATIVE, "synthetic"), Backend()
    ).draft(field, section, [item], decision)
    assert draft.status is FieldDraftStatus.INVALID


def test_chinese_heading_split_placeholder_format_and_citation_appendix(tmp_path: Path):
    source = tmp_path / "template.docx"
    output = tmp_path / "draft.docx"
    document = Document()
    style = document.styles.add_style("标题 1", WD_STYLE_TYPE.PARAGRAPH)
    document.add_paragraph("变更内容", style=style)
    paragraph = document.add_paragraph()
    paragraph.add_run("前缀 ")
    first = paragraph.add_run("{{变")
    first.bold = True
    paragraph.add_run("更内容}}")
    suffix = paragraph.add_run(" 后缀")
    suffix.italic = True
    document.save(source)

    schema = TemplateParser().parse(source)
    assert schema.sections[0].level == 1
    item = evidence("最大并发由500提升到1000。")
    from app.document_workflow.models import FieldDraft, SectionDraft, TaskStatus

    draft = SectionDraft(
        schema.sections[0].section_id,
        schema.sections[0].title,
        [FieldDraft(schema.fields[0].field_id, item.content, [item.evidence_id], FieldDraftStatus.DRAFTED, None, DraftingMode.EXTRACTIVE)],
        TaskStatus.COMPLETE,
    )
    TemplateRenderer().render(source, output, schema, [draft], [item])
    rendered = Document(output)
    body = rendered.paragraphs[1]
    assert "{{" not in body.text and "1000" in body.text
    assert any(run.bold and "1000" in run.text for run in body.runs)
    assert any(run.italic and "后缀" in run.text for run in body.runs)
    assert "引用来源" in [paragraph.text for paragraph in rendered.paragraphs]
    assert any("需求规格说明书" in paragraph.text and "V2.0" in paragraph.text for paragraph in rendered.paragraphs)

