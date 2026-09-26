from __future__ import annotations

import hashlib
from pathlib import Path

import pytest
from docx import Document

from app.document_workflow import DocumentWorkflowFacade, WorkflowScope
from app.document_workflow.configuration import DocumentWorkflowConfig
from app.document_workflow.evidence_models import normalize_text
from app.document_workflow.integrity import OutputIntegrityError
from app.document_workflow.rag import DemoRAGClient
from app.document_workflow.state import CheckpointInvalid
from app.document_workflow.workflow import DocumentWorkflow


def make_template(path: Path, fields=("变更内容", "影响范围", "回退方案")) -> Path:
    document = Document()
    document.add_heading("需求变更影响分析报告", level=1)
    for field in fields:
        document.add_heading(field, level=2)
        document.add_paragraph("{{" + field + "}}")
    document.save(path)
    return path


class BusinessRAG:
    def __init__(self, *, active_req="REQ-001@2.0", interrupt_on: str | None = None):
        self.active_req = active_req
        self.interrupt_on = interrupt_on
        self.calls: list[tuple[str, dict]] = []

    def active_versions(self):
        return {"REQ-001": self.active_req, "DESIGN-001": "DESIGN-001@1.0", "TEST-001": "TEST-001@1.0"}

    def retrieve(self, query: str, top_k: int, scope: dict | None = None):
        self.calls.append((query, dict(scope or {})))
        if self.interrupt_on and self.interrupt_on in query:
            raise KeyboardInterrupt("synthetic interruption")
        rows = []
        if "变更内容" in query:
            rows = [
                ("REQ-001", self.active_req, "V3.0" if self.active_req.endswith("3.0") else "V2.0", "REQUIREMENT", 12, "变更内容", "最大并发由500提升到1000，并增加峰值保护。"),
                ("DESIGN-001", "DESIGN-001@1.0", "V1.0", "DESIGN", 8, "设计影响", "连接池和限流器需要按1000并发重新配置。"),
            ]
        elif "影响范围" in query:
            rows = [("TEST-001", "TEST-001@1.0", "V1.0", "TEST", 21, "影响范围", "影响范围包括接口网关、容量测试和上线监控。")]
        # 回退方案 deliberately returns no evidence.
        results = []
        for rank, (doc, version, label, doc_type, page, section, content) in enumerate(rows[:top_k], start=1):
            results.append({
                "rank": rank,
                "similarity": 0.9 - rank * 0.05,
                "chunk_id": f"{version}:{page}:{section}",
                "document_id": doc,
                "project_id": "P-001",
                "document_type": doc_type,
                "version_id": version,
                "version_label": label,
                "version_status": "ACTIVE",
                "section_id": section,
                "section_path": [section],
                "page_number": page,
                "content": content,
                "content_hash": hashlib.sha256(normalize_text(content).encode()).hexdigest(),
            })
        return {"query": query, "results": results}


def config(runtime: Path) -> DocumentWorkflowConfig:
    return DocumentWorkflowConfig(
        rag_base_url=None,
        output_root=runtime / "outputs",
        data_classification="synthetic",
    )


def test_safe_business_e2e_scope_multi_evidence_missing_review_and_official_output(tmp_path: Path):
    runtime = tmp_path / "runtime"
    template = make_template(tmp_path / "change.docx")
    rag = BusinessRAG()
    facade = DocumentWorkflowFacade(runtime, config=config(runtime), rag_client=rag)
    scope = {"project_ids": ["P-001"], "document_types": ["REQUIREMENT", "DESIGN", "TEST"], "active_only": True}
    record = facade.create_workflow(template, scope)
    result = facade.run_workflow(record)

    assert result["workflow_status"] == "REVIEW_REQUIRED"
    assert result["scope"] == WorkflowScope.from_dict(scope).to_dict()
    assert all(call_scope == result["scope"] for _, call_scope in rag.calls)
    by_title = {item["section_title"]: item for item in result["sections"]}
    assert len(by_title["变更内容"]["fields"][0]["evidence_ids"]) == 2
    assert by_title["回退方案"]["fields"][0]["status"] == "MISSING"
    assert Path(result["artifacts"]["draft"]).is_file()

    workflow_id = result["workflow_id"]
    result = facade.review_section(workflow_id, section_id=by_title["变更内容"]["section_id"], action="APPROVED", reviewer="reviewer-a")
    result = facade.review_section(
        workflow_id,
        section_id=by_title["影响范围"]["section_id"],
        action="APPROVED",
        reviewer="reviewer-a",
        comment="人工精炼表达",
        edited_fields={by_title["影响范围"]["fields"][0]["field_id"]: "影响接口网关、容量测试、上线监控和回归验证。"},
    )
    result = facade.review_section(
        workflow_id,
        section_id=by_title["回退方案"]["section_id"],
        action="REJECTED",
        reviewer="reviewer-a",
        comment="知识库缺少回退方案",
    )
    assert result["workflow_status"] == "REJECTED"
    with pytest.raises(OutputIntegrityError, match="APPROVED"):
        facade.finalize_document(workflow_id)

    result = facade.review_section(
        workflow_id,
        section_id=by_title["回退方案"]["section_id"],
        action="APPROVED",
        reviewer="reviewer-a",
        comment="由评审人补充",
        edited_fields={by_title["回退方案"]["fields"][0]["field_id"]: "出现容量异常时回退至上一稳定版本，并恢复原限流参数。"},
    )
    assert result["all_sections_approved"] is True
    approved = facade.finalize_document(workflow_id)
    assert approved["workflow_status"] == "APPROVED"
    assert Path(approved["artifacts"]["approved"]).is_file()
    assert facade.list_workflows()[0]["workflow_id"] == workflow_id


def test_checkpoint_resume_preserves_scope_and_skips_completed_section(tmp_path: Path):
    template = make_template(tmp_path / "resume.docx", ("变更内容", "影响范围"))
    output = tmp_path / "output"
    scope = WorkflowScope(project_ids=("P-001",), active_only=True)
    cfg = DocumentWorkflowConfig(output_root=tmp_path, checkpoint_root=output / "checkpoints", trace_root=output)
    first = BusinessRAG(interrupt_on="影响范围")
    with pytest.raises(KeyboardInterrupt):
        DocumentWorkflow(first, config=cfg, scope=scope).run(template, output)
    checkpoint = next((output / "checkpoints").glob("dw_*.json"))
    workflow_id = checkpoint.stem

    resumed_client = BusinessRAG()
    result = DocumentWorkflow(resumed_client, config=cfg, scope=scope).resume(workflow_id, output=output, template=template)
    assert result["state"].scope_fingerprint == scope.fingerprint()
    assert not any("变更内容" in query for query, _ in resumed_client.calls)
    assert any("影响范围" in query for query, _ in resumed_client.calls)

    with pytest.raises(CheckpointInvalid, match="scope changed"):
        DocumentWorkflow(
            BusinessRAG(), config=cfg,
            scope=WorkflowScope(project_ids=("OTHER",), active_only=True),
        ).resume(workflow_id, output=output, template=template)


def test_stale_resume_refreshes_only_affected_section(tmp_path: Path):
    template = make_template(tmp_path / "stale.docx", ("变更内容", "影响范围"))
    output = tmp_path / "output"
    scope = WorkflowScope(project_ids=("P-001",), active_only=True)
    cfg = DocumentWorkflowConfig(output_root=tmp_path, checkpoint_root=output / "checkpoints", trace_root=output)
    initial = DocumentWorkflow(BusinessRAG(active_req="REQ-001@2.0"), config=cfg, scope=scope).run(template, output)
    workflow_id = initial["state"].workflow_id

    newer = BusinessRAG(active_req="REQ-001@3.0")
    resumed = DocumentWorkflow(newer, config=cfg, scope=scope).resume(workflow_id, output=output, template=template)
    assert any("变更内容" in query for query, _ in newer.calls)
    assert not any("影响范围" in query for query, _ in newer.calls)
    freshness = resumed["trace"]["freshness_events"]
    assert any(item["freshness"] == "STALE" for item in freshness)
    refresh = resumed["trace"]["refresh_summary"]
    assert refresh["total_sections"] == len(resumed["schema"].sections)
    assert len(refresh["affected_section_ids"]) == 1
    assert len(refresh["reused_section_ids"]) == 1
    assert len(refresh["re_retrieved_section_ids"]) == 1
    assert len(refresh["regenerated_section_ids"]) == 1
    assert len(refresh["unchanged_section_ids"]) == refresh["total_sections"] - len(refresh["affected_section_ids"])


def test_stale_evidence_blocks_official_output(tmp_path: Path):
    runtime = tmp_path / "runtime"
    template = make_template(tmp_path / "stale-final.docx", ("变更内容",))
    facade = DocumentWorkflowFacade(
        runtime,
        config=config(runtime),
        rag_client=BusinessRAG(),
    )
    record = facade.create_workflow(
        template, {"project_ids": ["P-001"], "active_only": True}
    )
    result = facade.run_workflow(record)
    section = next(item for item in result["sections"] if item["fields"])
    facade.review_section(
        result["workflow_id"],
        section_id=section["section_id"],
        action="APPROVED",
        reviewer="reviewer-a",
    )
    state = facade.repository.get(result["workflow_id"])
    state.evidence_snapshot["evidence"][0]["freshness"] = "STALE"
    facade.repository.save(state)

    with pytest.raises(OutputIntegrityError, match="EVIDENCE_(MEMBERSHIP|FRESHNESS)_INVALID"):
        facade.finalize_document(result["workflow_id"])


def test_demo_client_respects_document_scope():
    client = DemoRAGClient()
    payload = client.retrieve("变更内容", 5, {"document_ids": ["REQ-001"], "active_only": True})
    assert payload["results"]
    assert {item["document_id"] for item in payload["results"]} == {"REQ-001"}

