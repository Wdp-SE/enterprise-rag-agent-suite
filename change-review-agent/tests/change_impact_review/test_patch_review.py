from __future__ import annotations

import hashlib
import json
from datetime import datetime, timezone
from pathlib import Path

from docx import Document

from app.change_impact_review import (
    ChangeImpactCheckpointStore,
    ChangeImpactReviewFacade,
    ChangeImpactTaskState,
    ChangeTaskStatus,
    ParagraphPatchApplier,
    PatchApplyStatus,
    PatchCandidate,
    PatchExecutionService,
    PatchOperation,
    PatchReviewAction,
    PatchReviewService,
    PatchReviewStatus,
)
from app.document_workflow.evidence_models import Evidence
from app.document_workflow.state import CheckpointStore


def docx(path: Path, body: str) -> Path:
    document = Document()
    document.add_heading("系统设计", level=1)
    document.add_paragraph(body)
    document.add_paragraph("未受影响的段落")
    document.save(path)
    return path


def evidence(*, version: str = "design-v1", content: str = "REQ-023 并发提升到 1000") -> Evidence:
    return Evidence(
        title="需求规格说明书",
        content=content,
        organization="星海软件科技有限公司（虚构）",
        source_type="RAG",
        document_id="requirements",
        project_id="PAYMENT",
        document_type="REQUIREMENT",
        version_id=version,
        version_label="V2.0",
        version_status="ACTIVE",
        chunk_id="requirements-v2:REQ-023",
        query="REQ-023",
        section="REQUIREMENTS",
        section_path=["项目需求", "REQUIREMENTS"],
        page_number=1,
        retrieved_at=datetime.now(timezone.utc),
    )


def patch_for(original: str, ev: Evidence) -> PatchCandidate:
    assert ev.evidence_id and ev.content_hash
    return PatchCandidate.create(
        target_document_id="system_design",
        base_version_id="design-v1",
        target_section_id="DESIGN",
        target_anchor="paragraph:1",
        operation=PatchOperation.REPLACE_PARAGRAPH,
        original_content=original,
        proposed_content="DES-014 面向 REQ-023：连接池支持 1000 并发，并通过异步任务状态返回结果。",
        reason="需求并发与接口模式发生变化",
        evidence_ids=[ev.evidence_id],
        evidence_content_hashes={ev.evidence_id: ev.content_hash},
        scope_fingerprint="scope-1",
    )


def test_patch_requires_approval_and_edit_requires_reconfirmation(tmp_path: Path) -> None:
    original = "DES-014 面向 REQ-023：连接池支持 500 并发，并同步返回结果。"
    source = docx(tmp_path / "design.docx", original)
    ev = evidence()
    patch = patch_for(original, ev)
    applier = ParagraphPatchApplier()

    blocked = applier.apply(
        patch,
        source=source,
        output=tmp_path / "blocked.docx",
        current_version_id="design-v1",
        applied_keys=set(),
    )
    assert blocked.status is PatchApplyStatus.BLOCKED_REVIEW
    assert not (tmp_path / "blocked.docx").exists()

    review = PatchReviewService()
    edited = review.review(
        patch,
        action=PatchReviewAction.EDIT,
        reviewer="reviewer-a",
        edited_content="DES-014 人工修订：支持 1000 并发和异步状态查询。",
        comment="按设计口径调整",
    )
    assert edited.status is PatchReviewStatus.EDITED_RECONFIRM_REQUIRED
    assert patch.review_status is PatchReviewStatus.EDITED_RECONFIRM_REQUIRED
    assert applier.apply(
        patch,
        source=source,
        output=tmp_path / "still-blocked.docx",
        current_version_id="design-v1",
        applied_keys=set(),
    ).status is PatchApplyStatus.BLOCKED_REVIEW

    approved = review.review(
        patch,
        action=PatchReviewAction.APPROVE,
        reviewer="reviewer-a",
        comment="复核通过",
    )
    assert approved.status is PatchReviewStatus.APPROVED


def test_approved_patch_changes_only_target_and_is_idempotent(tmp_path: Path) -> None:
    original = "DES-014 面向 REQ-023：连接池支持 500 并发，并同步返回结果。"
    source = docx(tmp_path / "design.docx", original)
    source_hash = hashlib.sha256(source.read_bytes()).hexdigest()
    ev = evidence()
    patch = patch_for(original, ev)
    PatchReviewService().review(patch, action=PatchReviewAction.APPROVE, reviewer="reviewer-a")
    applied_keys: set[str] = set()
    output = tmp_path / "candidate.docx"

    first = ParagraphPatchApplier().apply(
        patch,
        source=source,
        output=output,
        current_version_id="design-v1",
        applied_keys=applied_keys,
    )
    second = ParagraphPatchApplier().apply(
        patch,
        source=source,
        output=output,
        current_version_id="design-v1",
        applied_keys=applied_keys,
    )

    assert first.status is PatchApplyStatus.APPLIED
    assert second.status is PatchApplyStatus.SKIPPED_ALREADY_APPLIED
    assert hashlib.sha256(source.read_bytes()).hexdigest() == source_hash
    rendered = Document(output)
    assert rendered.paragraphs[1].text == patch.proposed_content
    assert rendered.paragraphs[2].text == "未受影响的段落"
    assert len(applied_keys) == 1


def test_external_target_change_returns_conflict_without_output(tmp_path: Path) -> None:
    original = "DES-014 面向 REQ-023：连接池支持 500 并发，并同步返回结果。"
    changed = docx(tmp_path / "changed.docx", original + " 外部修改")
    ev = evidence()
    patch = patch_for(original, ev)
    PatchReviewService().review(patch, action=PatchReviewAction.APPROVE, reviewer="reviewer-a")

    result = ParagraphPatchApplier().apply(
        patch,
        source=changed,
        output=tmp_path / "candidate.docx",
        current_version_id="design-v1",
        applied_keys=set(),
    )

    assert result.status is PatchApplyStatus.CONFLICT
    assert result.failure_reason == "ORIGINAL_CONTENT_CHANGED"
    assert not (tmp_path / "candidate.docx").exists()


def test_stale_evidence_blocks_only_dependent_patch(tmp_path: Path) -> None:
    original = "DES-014 面向 REQ-023：连接池支持 500 并发，并同步返回结果。"
    source = docx(tmp_path / "design.docx", original)
    ev = evidence(version="requirements-v1")
    patch = patch_for(original, ev)
    PatchReviewService().review(patch, action=PatchReviewAction.APPROVE, reviewer="reviewer-a")

    result = PatchExecutionService(
        active_version_for_document=lambda document_id: "requirements-v2"
    ).apply(
        patch,
        evidence_by_id={ev.evidence_id: ev},
        expected_scope_fingerprint="scope-1",
        source=source,
        output=tmp_path / "candidate.docx",
        current_version_id="design-v1",
        applied_keys=set(),
    )

    assert result.status is PatchApplyStatus.BLOCKED_STALE_EVIDENCE
    assert result.failure_reason == "EVIDENCE_STALE"
    assert not (tmp_path / "candidate.docx").exists()


def test_change_task_checkpoint_reuses_atomic_store_and_preserves_idempotency(tmp_path: Path) -> None:
    ev = evidence()
    patch = patch_for("original", ev)
    state = ChangeImpactTaskState(
        task_id="cir_1234567890abcdef",
        organization_id="demo_company_a",
        project_id="PAYMENT",
        scope={"project_ids": ["PAYMENT"], "active_only": True},
        scope_fingerprint="scope-1",
        base_document_id="system_design",
        base_version_id="design-v1",
        source_path=str(tmp_path / "source.docx"),
        candidate_path=str(tmp_path / "candidate.docx"),
        status=ChangeTaskStatus.REVIEW_REQUIRED,
        patches=[patch],
        applied_keys={patch.idempotency_key},
    )
    store = ChangeImpactCheckpointStore(CheckpointStore(tmp_path / "checkpoints"))

    store.save(state)
    resumed = store.load(state.task_id)

    assert resumed.task_id == state.task_id
    assert resumed.status is ChangeTaskStatus.REVIEW_REQUIRED
    assert resumed.applied_keys == {patch.idempotency_key}
    assert resumed.patches[0].patch_id == patch.patch_id



def test_facade_checkpoint_resume_does_not_apply_patch_twice(tmp_path: Path) -> None:
    original = "DES-014 面向 REQ-023：连接池支持 500 并发，并同步返回结果。"
    source = docx(tmp_path / "design.docx", original)
    ev = evidence(version="requirements-v2")
    patch = patch_for(original, ev)
    scope = {"project_ids": ["PAYMENT"], "active_only": True}
    patch.scope_fingerprint = hashlib.sha256(
        json.dumps(scope, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode("utf-8")
    ).hexdigest()
    facade = ChangeImpactReviewFacade(
        tmp_path / "runtime",
        active_version_for_document=lambda document_id: "requirements-v2",
    )
    task = facade.create_task(
        organization_id="demo_company_a",
        project_id="PAYMENT",
        scope=scope,
        base_document_id="system_design",
        base_version_id="design-v1",
        source=source,
        impacts=[{
            "changed_item_id": "req-item",
            "impacted_item_id": "design-item",
            "discovery_source": "EXPLICIT_TRACE",
            "evidence_ids": [ev.evidence_id],
            "trace_link_id": "trace-1",
            "review_status": "CONFIRMED",
        }],
        patches=[patch],
        evidence=[ev],
    )
    reviewed = facade.review_patch(
        task.task_id,
        patch.patch_id,
        action=PatchReviewAction.APPROVE,
        reviewer="reviewer-a",
    )
    assert reviewed.status is ChangeTaskStatus.APPLY_READY

    first, first_results = facade.apply_reviewed_patches(
        task.task_id, current_version_id="design-v1"
    )
    first_hash = hashlib.sha256(Path(first.candidate_path).read_bytes()).hexdigest()
    resumed, resumed_results = facade.apply_reviewed_patches(
        task.task_id, current_version_id="design-v1"
    )

    assert first.status is ChangeTaskStatus.CANDIDATE_READY
    assert first_results[0].status is PatchApplyStatus.APPLIED
    assert resumed.status is ChangeTaskStatus.CANDIDATE_READY
    assert resumed_results[0].status is PatchApplyStatus.SKIPPED_ALREADY_APPLIED
    assert hashlib.sha256(Path(resumed.candidate_path).read_bytes()).hexdigest() == first_hash
    assert len(resumed.applied_keys) == 1
    assert any(row["step"] == "APPLY_PATCH" for row in resumed.trace)
