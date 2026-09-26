from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path

from docx import Document

from app.change_impact_review import (
    PatchCandidate,
    PatchExecutionService,
    PatchOperation,
    PatchReviewAction,
    PatchReviewService,
    QualityGateReason,
    QualityGateStatus,
)
from app.document_workflow.evidence_models import Evidence


def write_docx(path: Path, body: str) -> Path:
    document = Document()
    document.add_heading("系统设计", level=1)
    document.add_paragraph(body)
    document.save(path)
    return path


def evidence(*, version: str = "requirements-v2", project_id: str = "PAYMENT") -> Evidence:
    return Evidence(
        title="需求规格说明书",
        content="REQ-023 最大并发提升到 1000。",
        organization="星海软件科技有限公司（虚构）",
        source_type="RAG",
        document_id="requirements",
        project_id=project_id,
        document_type="REQUIREMENT",
        version_id=version,
        version_label="V2.0",
        version_status="ACTIVE",
        freshness="FRESH",
        chunk_id=f"{version}:REQ-023",
        section="性能要求",
        section_path=["需求规格", "性能要求"],
        page_number=1,
        retrieved_at=datetime.now(timezone.utc),
    )


def patch(original: str, item: Evidence) -> PatchCandidate:
    assert item.evidence_id and item.content_hash
    return PatchCandidate.create(
        target_document_id="system_design",
        base_version_id="design-v1",
        target_section_id="DESIGN",
        target_anchor="paragraph:1",
        operation=PatchOperation.REPLACE_PARAGRAPH,
        original_content=original,
        proposed_content="DES-014：连接池支持 1000 并发。",
        reason="需求并发上限变化",
        evidence_ids=[item.evidence_id],
        evidence_content_hashes={item.evidence_id: item.content_hash},
        scope_fingerprint="scope-1",
    )


def evaluate(
    service: PatchExecutionService,
    candidate: PatchCandidate,
    item: Evidence,
    source: Path,
    **overrides,
):
    values = {
        "evidence_by_id": {item.evidence_id: item},
        "expected_scope_fingerprint": "scope-1",
        "expected_scope": {"project_ids": ["PAYMENT"], "active_only": True},
        "source": source,
        "output": source.with_name("candidate.docx"),
        "current_version_id": "design-v1",
        "applied_keys": set(),
    }
    values.update(overrides)
    return service.evaluate(candidate, **values)


def test_quality_gate_passes_only_after_deterministic_checks_and_review(tmp_path: Path) -> None:
    original = "DES-014：连接池支持 500 并发。"
    source = write_docx(tmp_path / "design.docx", original)
    item = evidence()
    candidate = patch(original, item)
    service = PatchExecutionService(lambda document_id: "requirements-v2")

    before_review = evaluate(service, candidate, item, source)
    assert before_review.status is QualityGateStatus.BLOCKED
    assert before_review.reasons == [QualityGateReason.REVIEW_REQUIRED]

    PatchReviewService().review(
        candidate,
        action=PatchReviewAction.APPROVE,
        reviewer="reviewer-a",
    )
    passed = evaluate(service, candidate, item, source)

    assert passed.status is QualityGateStatus.PASS
    assert passed.reasons == []


def test_quality_gate_reports_no_evidence_scope_stale_and_base_version(tmp_path: Path) -> None:
    original = "DES-014：连接池支持 500 并发。"
    source = write_docx(tmp_path / "design.docx", original)
    item = evidence()
    candidate = patch(original, item)
    PatchReviewService().review(candidate, action=PatchReviewAction.APPROVE, reviewer="reviewer-a")
    service = PatchExecutionService(lambda document_id: "requirements-v2")

    missing = evaluate(service, candidate, item, source, evidence_by_id={})
    assert missing.reasons == [QualityGateReason.NO_EVIDENCE]

    candidate.scope_fingerprint = "different"
    outside = evaluate(service, candidate, item, source)
    assert outside.reasons == [QualityGateReason.OUT_OF_SCOPE_EVIDENCE]
    candidate.scope_fingerprint = "scope-1"

    outside_item = evidence(project_id="OTHER")
    outside_candidate = patch(original, outside_item)
    PatchReviewService().review(
        outside_candidate,
        action=PatchReviewAction.APPROVE,
        reviewer="reviewer-a",
    )
    outside_evidence = evaluate(
        service,
        outside_candidate,
        outside_item,
        source,
    )
    assert outside_evidence.reasons == [QualityGateReason.OUT_OF_SCOPE_EVIDENCE]

    stale_service = PatchExecutionService(lambda document_id: "requirements-v3")
    stale = evaluate(stale_service, candidate, item, source)
    assert stale.reasons == [QualityGateReason.STALE_EVIDENCE]

    changed_base = evaluate(service, candidate, item, source, current_version_id="design-v2")
    assert changed_base.reasons == [QualityGateReason.BASE_VERSION_MISMATCH]


def test_quality_gate_reports_target_change_and_duplicate_application(tmp_path: Path) -> None:
    original = "DES-014：连接池支持 500 并发。"
    source = write_docx(tmp_path / "design.docx", original + " 外部改动")
    item = evidence()
    candidate = patch(original, item)
    PatchReviewService().review(candidate, action=PatchReviewAction.APPROVE, reviewer="reviewer-a")
    service = PatchExecutionService(lambda document_id: "requirements-v2")

    changed = evaluate(service, candidate, item, source)
    assert changed.reasons == [QualityGateReason.TARGET_CHANGED]

    duplicate = evaluate(
        service,
        candidate,
        item,
        source,
        applied_keys={candidate.idempotency_key},
    )
    assert duplicate.reasons == [QualityGateReason.ALREADY_APPLIED]
