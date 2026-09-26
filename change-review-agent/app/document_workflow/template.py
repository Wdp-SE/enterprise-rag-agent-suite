"""Parse and render the supported subset of structured DOCX templates."""

from __future__ import annotations

import hashlib
import re
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

from .evidence_models import Evidence
from .models import SectionDraft, TemplateField, TemplateSchema, TemplateSection, TemplateTable


PLACEHOLDER = re.compile(r"\{\{\s*([a-zA-Z0-9_\u4e00-\u9fff]+)\s*\}\}|【待填写】")
HEADING = re.compile(r"^(?:Heading|标题)\s*([123])$", re.IGNORECASE)


def body_blocks(document):
    for child in document.element.body.iterchildren():
        if child.tag.endswith("}p"):
            yield Paragraph(child, document)
        elif child.tag.endswith("}tbl"):
            yield Table(child, document)


def _heading_level(paragraph: Paragraph) -> int | None:
    match = HEADING.match(paragraph.style.name or "")
    if match:
        return int(match.group(1))
    properties = paragraph._p.pPr
    outline = properties.find(qn("w:outlineLvl")) if properties is not None else None
    if outline is not None:
        value = int(outline.get(qn("w:val"))) + 1
        return value if value in {1, 2, 3} else None
    return None


def _replace_paragraph(paragraph: Paragraph, old: str, new: str) -> None:
    if not old:
        return
    full = "".join(run.text for run in paragraph.runs)
    start = full.find(old)
    if start < 0:
        return
    end = start + len(old)
    cursor = 0
    inserted = False
    for run in paragraph.runs:
        run_start, run_end = cursor, cursor + len(run.text)
        cursor = run_end
        if run_end <= start or run_start >= end:
            continue
        prefix = run.text[: max(0, start - run_start)] if run_start <= start < run_end else ""
        suffix = run.text[max(0, end - run_start):] if run_start < end <= run_end else ""
        if not inserted:
            run.text = prefix + new + suffix
            inserted = True
        else:
            run.text = suffix


class TemplateParser:
    def parse(self, path: str | Path) -> TemplateSchema:
        source = Path(path).resolve()
        if source.suffix.lower() != ".docx" or not source.is_file():
            raise ValueError("template must be an existing .docx file")
        document = Document(source)
        sections: list[TemplateSection] = []
        fields: list[TemplateField] = []
        tables: list[TemplateTable] = []
        stack: list[TemplateSection] = []
        for block_index, block in enumerate(body_blocks(document)):
            if isinstance(block, Paragraph):
                level = _heading_level(block)
                if level is not None:
                    while stack and stack[-1].level >= level:
                        stack.pop()
                    section = TemplateSection(
                        f"s{len(sections)+1:03d}", block.text.strip(), level, len(sections)+1,
                        stack[-1].section_id if stack else None, block_index,
                    )
                    sections.append(section)
                    stack.append(section)
                elif stack:
                    for match in PLACEHOLDER.finditer(block.text):
                        fields.append(TemplateField(
                            f"f{len(fields)+1:03d}", match.group(1) or stack[-1].title,
                            "paragraph", True, match.group(0), {"block": block_index}, stack[-1].section_id,
                        ))
            elif isinstance(block, Table):
                table = TemplateTable(f"t{len(tables)+1:03d}", len(block.rows))
                if stack:
                    for row_index, row in enumerate(block.rows):
                        for cell_index, cell in enumerate(row.cells):
                            for match in PLACEHOLDER.finditer(cell.text):
                                label = row.cells[cell_index - 1].text.strip() if cell_index else ""
                                location = {"block": block_index, "row": row_index, "cell": cell_index}
                                table.fillable_cells.append(location)
                                fields.append(TemplateField(
                                    f"f{len(fields)+1:03d}", match.group(1) or label or stack[-1].title,
                                    "table", True, match.group(0), location, stack[-1].section_id,
                                ))
                tables.append(table)
        if not sections:
            raise ValueError("template has no Heading/标题/Outline Level 1-3 sections")
        explicit = {item.section_id for item in fields}
        for section in sections:
            if section.section_id not in explicit and not any(child.parent_section_id == section.section_id for child in sections):
                fields.append(TemplateField(
                    f"f{len(fields)+1:03d}", section.title, "implicit", True, None,
                    {"block": section.heading_location}, section.section_id,
                ))
        digest = hashlib.sha256(source.read_bytes()).hexdigest()[:16]
        return TemplateSchema(f"tpl_{digest}", source.name, sections, fields, tables)


class TemplateRenderer:
    def render(
        self,
        template: str | Path,
        output: str | Path,
        schema: TemplateSchema,
        drafts: list[SectionDraft],
        evidence: list[Evidence] = (),
    ) -> Path:
        source, target = Path(template).resolve(), Path(output).resolve()
        if source == target:
            raise ValueError("draft output must differ from original template")
        document = Document(source)
        blocks = list(body_blocks(document))
        field_drafts = {item.field_id: item for draft in drafts for item in draft.fields}
        used_ids = sorted({eid for item in field_drafts.values() for eid in item.evidence_ids})
        citation_numbers = {evidence_id: index for index, evidence_id in enumerate(used_ids, start=1)}
        for field in schema.fields:
            draft = field_drafts.get(field.field_id)
            value = draft.content if draft is not None else "[MISSING: 当前知识库中未发现相关证据]"
            if draft is not None and draft.evidence_ids:
                value += " " + "".join(f"[{citation_numbers[item]}]" for item in draft.evidence_ids if item in citation_numbers)
            block = blocks[field.location["block"]]
            if field.field_type == "table":
                cell = block.cell(field.location["row"], field.location["cell"])
                for paragraph in cell.paragraphs:
                    _replace_paragraph(paragraph, field.placeholder or "", value)
            elif field.field_type == "paragraph":
                _replace_paragraph(block, field.placeholder or "", value)
            else:
                element = OxmlElement("w:p")
                block._p.addnext(element)
                Paragraph(element, block._parent).add_run(value)
        if used_ids:
            document.add_heading("引用来源", level=1)
            evidence_by_id = {item.evidence_id: item for item in evidence if item.evidence_id}
            for evidence_id in used_ids:
                item = evidence_by_id.get(evidence_id)
                if item is None:
                    raise ValueError("citation references unknown Evidence")
                section = " / ".join(item.section_path or []) or item.section or "—"
                document.add_paragraph(
                    f"[{citation_numbers[evidence_id]}] 文档：{item.title}；版本：{item.version_label or item.version_id or '—'}；"
                    f"状态：{item.version_status or '—'}；章节：{section}；页码：{item.page_number or '—'}；Evidence ID：{evidence_id}"
                )
        target.parent.mkdir(parents=True, exist_ok=True)
        document.save(target)
        return target
