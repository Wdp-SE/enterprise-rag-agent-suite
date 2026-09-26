"""Create the retained requirement-change impact analysis template."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt


def create(path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("需求变更影响分析报告")
    run.bold = True
    run.font.size = Pt(20)

    document.add_heading("1 变更概述", level=1)
    document.add_paragraph("{{变更内容}}")

    document.add_heading("2 影响分析", level=1)
    document.add_paragraph("{{影响范围}}")

    document.add_heading("3 验证结论", level=1)
    table = document.add_table(rows=2, cols=2)
    table.style = "Table Grid"
    table.cell(0, 0).text = "验证项目"
    table.cell(0, 1).text = "结论"
    table.cell(1, 0).text = "容量与稳定性"
    table.cell(1, 1).text = "{{验证结果}}"

    document.add_heading("4 回退方案", level=1)
    document.add_paragraph("{{回退方案}}")
    document.save(path)
    return path


if __name__ == "__main__":
    output = Path("project_delivery/document_workflow_business_refactor/demo/requirement_change_impact_template.docx")
    print(create(output))

